"""Batches in Excel: the Batch column on the Students sheet (download and upload), the Batches
sheet of Download everything, and names the upload doesn't know."""

from __future__ import annotations

import datetime as dt
import io

from fastapi.testclient import TestClient
from helpers import make_student
from openpyxl import load_workbook
from test_excel_import import commit, commit_body, preview, xlsx

HEAD = ["Name", "Monthly fee", "Joined", "Batch"]


def make_batch(api: TestClient, **fields: object) -> dict:
    response = api.post("/api/batches", json=fields)
    assert response.status_code == 201, response.text
    return response.json()


def students(api: TestClient) -> dict[str, tuple[str | None, str | None]]:
    rows = api.get("/api/students", params={"status": "all"}).json()
    return {s["name"]: (s["batch_name"], s["batch_label"]) for s in rows}


def test_the_batch_column_finds_batches_by_name(api: TestClient) -> None:
    make_batch(api, name="Mon/Wed Evening")
    data = xlsx(
        ("Students", [HEAD,
                      ["Ananya Rao", 1500, "Jun 2026", "mon/wed  EVENING"],
                      ["Kabir Mehta", 1500, "Jun 2026", None]]),
    )  # fmt: skip
    shown = preview(api, data)
    [batch] = shown["batches"]
    assert (batch["name"], batch["status"], batch["student_count"]) == (
        "Mon/Wed Evening",
        "exists",
        1,
    )
    assert {s["name"]: s["batch_name"] for s in shown["students"]} == {
        "Ananya Rao": "mon/wed  EVENING",
        "Kabir Mehta": None,
    }
    result = commit(api, shown)
    assert (result["students_added"], result["batches_added"]) == (2, 0)
    assert students(api) == {"Ananya Rao": ("Mon/Wed Evening", None), "Kabir Mehta": (None, None)}


def test_an_unknown_batch_is_flagged_and_never_created_by_itself(api: TestClient) -> None:
    data = xlsx(
        ("Students", [["Name", "Monthly fee", "Class/batch"],
                      ["Ananya Rao", 1500, "Thu 7pm Adults"],
                      ["Kabir Mehta", 1500, "thu 7PM adults"]]),
    )  # fmt: skip
    shown = preview(api, data)
    [batch] = shown["batches"]
    assert batch == {
        "name": "Thu 7pm Adults",
        "status": "not_found",
        "reason": "Batch not found, will be left without a batch",
        "row": None,
        "student_count": 2,
        "batch_id": None,
    }
    commit(api, shown)
    assert api.get("/api/batches").json() == []
    # Left without a batch; what the file said is kept as their old class label.
    assert students(api) == {
        "Ananya Rao": (None, "Thu 7pm Adults"),
        "Kabir Mehta": (None, "thu 7PM adults"),
    }


def test_an_unknown_batch_is_created_when_chosen(api: TestClient) -> None:
    data = xlsx(("Students", [HEAD, ["Ananya Rao", 1500, "Jun 2026", "Thu 7pm Adults"]]))
    shown = preview(api, data)
    body = commit_body(shown)
    body["create_batches"] = ["thu 7pm adults"]
    response = api.post("/api/import/commit", json=body)
    assert response.status_code == 200, response.text
    assert response.json()["batches_added"] == 1
    [batch] = api.get("/api/batches").json()
    assert (batch["name"], batch["days"], batch["default_fee_paise"]) == (
        "Thu 7pm Adults",
        [],
        None,
    )
    assert students(api) == {"Ananya Rao": ("Thu 7pm Adults", None)}


def test_only_batches_the_preview_offered_can_be_created(api: TestClient) -> None:
    make_batch(api, name="Saturday")
    data = xlsx(("Students", [HEAD, ["Ananya Rao", 1500, "Jun 2026", "Saturday"]]))
    shown = preview(api, data)
    for name in ("Saturday", "Something else"):
        body = commit_body(shown)
        body["create_batches"] = [name]
        response = api.post("/api/import/commit", json=body)
        assert response.status_code == 422
        assert response.json()["detail"][0]["loc"] == ["body", "create_batches"]
    assert len(api.get("/api/batches").json()) == 1
    assert api.get("/api/students").json() == []


def test_a_batches_sheet_adds_its_batches_with_their_details(api: TestClient) -> None:
    data = xlsx(
        ("Students", [HEAD, ["Ananya Rao", 1500, "Jun 2026", "Evening"]]),
        ("Batches", [["Name", "Location", "Days", "Starts", "Ends", "Usual monthly fee ₹", "Notes"],
                     ["Evening", "Koramangala", "Mon, Wed", "5:00 pm", dt.time(18, 0), 1800, "B"],
                     ["Weekend", None, "Sat & Sun", "9.30 am", "11:00", None, None],
                     ["Broken", None, "Someday", None, None, None, None],
                     ["evening", None, None, None, None, None, None]]),
    )  # fmt: skip
    shown = preview(api, data)
    by_name = {b["name"]: b for b in shown["batches"]}
    assert (by_name["Evening"]["status"], by_name["Evening"]["student_count"]) == ("new", 1)
    assert by_name["Weekend"]["status"] == "new"
    assert by_name["Broken"]["status"] == "problem"
    assert "isn't a list of days" in by_name["Broken"]["reason"]
    twice = [b for b in shown["batches"] if b["row"] == 5]
    assert twice[0]["status"] == "problem" and "Listed twice" in twice[0]["reason"]

    result = commit(api, shown)
    assert result["batches_added"] == 2
    made = {b["name"]: b for b in api.get("/api/batches").json()}
    assert {k: made["Evening"][k] for k in ("location", "days", "start_time", "end_time",
                                            "default_fee_paise", "notes")} == {
        "location": "Koramangala",
        "days": ["mon", "wed"],
        "start_time": "17:00",
        "end_time": "18:00",
        "default_fee_paise": 180000,
        "notes": "B",
    }  # fmt: skip
    assert (made["Weekend"]["days"], made["Weekend"]["start_time"]) == (["sat", "sun"], "09:30")
    assert students(api) == {"Ananya Rao": ("Evening", None)}


def test_existing_students_are_never_moved(api: TestClient) -> None:
    make_batch(api, name="Saturday")
    make_student(api, name="Ananya Rao", joined_month="2026-06")
    data = xlsx(("Students", [HEAD, ["Ananya Rao", 1500, "Jun 2026", "Saturday"]]))
    shown = preview(api, data)
    assert shown["students"][0]["status"] == "exists"
    commit(api, shown)
    assert students(api) == {"Ananya Rao": (None, None)}


def test_the_students_download_has_the_batch_and_follows_the_tab(api: TestClient) -> None:
    sat = make_batch(api, name="Saturday")
    make_student(api, name="Ananya Rao", batch_id=sat["id"], batch_label="Sat 10am")
    make_student(api, name="Kabir Mehta")

    def rows(**params: str) -> list[list[object]]:
        response = api.get("/api/export/students.xlsx", params=params)
        assert response.status_code == 200
        ws = load_workbook(io.BytesIO(response.content)).active
        return [[c.value for c in r][:5] for r in ws.iter_rows()]

    assert rows(batch=str(sat["id"])) == [
        ["Name", "Phone", "Parent/guardian", "Batch", "Old class label"],
        ["Ananya Rao", None, None, "Saturday", "Sat 10am"],
    ]
    assert [r[0] for r in rows(batch="none")[1:]] == ["Kabir Mehta"]
    assert [r[0] for r in rows(q="saturday")[1:]] == ["Ananya Rao"]  # the search finds batches
    assert api.get("/api/export/students.xlsx", params={"batch": "x"}).status_code == 422


def test_the_template_has_a_batch_column(api: TestClient) -> None:
    response = api.get("/api/import/template.xlsx", params={"kind": "students"})
    ws = load_workbook(io.BytesIO(response.content))["Students"]
    assert "Batch" in [c.value for c in ws[1]]


def test_the_report_download_filters_and_groups_by_batch(api: TestClient) -> None:
    sat = make_batch(api, name="Saturday")
    b10 = make_batch(api, name="Batch 10")
    b2 = make_batch(api, name="Batch 2")
    make_student(api, name="Ananya Rao", batch_id=sat["id"], joined_month="2026-06")
    make_student(api, name="Kabir Mehta", batch_id=b10["id"], joined_month="2026-06")
    make_student(api, name="Meera Iyer", batch_id=b2["id"], joined_month="2026-06")
    make_student(api, name="Diya Nair", joined_month="2026-06")

    def sheet_rows(**params: str) -> list[list[object]]:
        response = api.get("/api/report.xlsx", params={"month": "2026-06", **params})
        assert response.status_code == 200, response.text
        ws = load_workbook(io.BytesIO(response.content)).active
        return [[c.value for c in r] for r in ws.iter_rows()]

    only_sat = sheet_rows(batch=str(sat["id"]))
    assert "in Saturday" in only_sat[0][0]
    names = [r[0] for r in only_sat if r[0] in ("Ananya Rao", "Kabir Mehta", "Meera Iyer")]
    assert names == ["Ananya Rao"]
    assert [r[0] for r in sheet_rows(batch="none") if r[0] == "Diya Nair"] == ["Diya Nair"]

    grouped = sheet_rows(group="batch")
    assert "grouped by batch" in grouped[0][0]
    firsts = [r[0] for r in grouped]
    headings = [f for f in firsts if isinstance(f, str) and f.endswith("(1 student)")]
    # Numbers in number order, "No batch" last; each heading right above its student.
    assert headings == [
        "Batch 2 (1 student)",
        "Batch 10 (1 student)",
        "Saturday (1 student)",
        "No batch (1 student)",
    ]
    assert firsts[firsts.index("Batch 2 (1 student)") + 1] == "Meera Iyer"
    assert firsts[firsts.index("No batch (1 student)") + 1] == "Diya Nair"
    assert api.get("/api/report.xlsx", params={"group": "class"}).status_code == 422
