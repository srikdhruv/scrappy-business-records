"""The monthly report (/api/report and /api/report.xlsx), current month frozen at June 2026.

The report must reconcile exactly with the dashboard and the profiles: `assert_reconciles`
checks every row and every total against them, for every month in a range, and runs on each
scenario below and on the demo data.
"""

from __future__ import annotations

import io
from typing import Any

import pytest
from conftest import FROZEN_TODAY
from fastapi.testclient import TestClient
from helpers import make_student, pay
from openpyxl import load_workbook

from app.db import session_factory
from app.seed import seed

Json = dict[str, Any]

MONTHS = [f"2025-{m:02d}" for m in range(10, 13)] + [f"2026-{m:02d}" for m in range(1, 13)]


def report(api: TestClient, month: str | None = None) -> Json:
    response = api.get("/api/report", params={"month": month} if month else {})
    assert response.status_code == 200, response.text
    return response.json()


def rows_by_name(body: Json) -> dict[str, Json]:
    return {r["student_name"]: r for r in body["rows"]}


def _relevant(student: Json, detail: Json, month: str, owes_earlier: bool) -> bool:
    enrolled = student["joined_month"] <= month and (
        student["left_month"] is None or month <= student["left_month"]
    )
    line = next((m for m in detail["months"] if m["month"] == month), None)
    money = line is not None and (line["paid_paise"] > 0 or line["covered_by_credit_paise"] > 0)
    return enrolled or money or owes_earlier


def assert_reconciles(api: TestClient, months: list[str] = MONTHS) -> None:
    """Every number on the report is the dashboard's or the profile's, for every month."""
    students = {s["id"]: s for s in api.get("/api/students", params={"status": "all"}).json()}
    details = {i: api.get(f"/api/students/{i}").json() for i in students}
    for month in months:
        body = report(api, month)
        board = api.get("/api/dashboard", params={"month": month}).json()
        rows = body["rows"]
        t = body["totals"]
        s = board["summary"]

        # Totals are the dashboard summary.
        assert t["fee_paise"] == s["expected_paise"], month
        assert t["collected_paise"] == s["collected_paise"], month
        assert t["short_paise"] == s["still_due_paise"], month
        assert t["not_fully_paid_count"] == s["not_fully_paid_count"], month
        assert t["active_student_count"] == s["active_student_count"], month
        if month > body["current_month"]:
            assert t["collected_paise"] == s["paid_ahead_paise"], month
        # ...and the sums of the rows.
        for field in (
            "fee_paise",
            "paid_paise",
            "paid_direct_paise",
            "covered_by_credit_paise",
            "extra_sent_paise",
            "extra_unused_paise",
            "short_paise",
            "owed_before_paise",
            "owed_now_paise",
            "credit_paise",
            "paid_ahead_paise",
        ):
            assert t[field] == sum(r[field] for r in rows), (month, field)
        assert t["student_count"] == len(rows)

        # Yet to pay and earlier months still owed are exactly the rows that say so.
        assert {r["student_id"] for r in rows if r["short_paise"] > 0} == {
            y["student_id"] for y in board["yet_to_pay"]
        }, month
        backlog = {b["student_id"]: b for b in board["backlog"]}
        assert {r["student_id"] for r in rows if r["owed_before_paise"] > 0} == set(backlog)
        for sid, b in backlog.items():
            row = next(r for r in rows if r["student_id"] == sid)
            assert row["owed_before_paise"] == b["total_owed_paise"], month
            assert row["owed_before_months"] == [m["month"] for m in b["months"]], month

        # Extra money into M, and out of the money logged for M, seen from the other months'
        # side on the profiles (built from the per-month ledger fields, never credit_moves).
        lines = [line for d in details.values() for line in d["months"]]
        assert t["covered_by_credit_paise"] == sum(
            e["amount_paise"]
            for line in lines
            for e in line["extra_sent"]
            if e["to_month"] == month
        ), month
        assert t["extra_sent_paise"] == sum(
            c["amount_paise"]
            for line in lines
            for c in line["credit_sources"]
            if c["for_month"] == month
        ), month

        # Every row is the profile's month and the students list's standing.
        for r in rows:
            student = students[r["student_id"]]
            assert r["student_name"] == student["name"]
            assert r["owed_now_paise"] == student["owed_paise"]
            assert r["credit_paise"] == student["credit_paise"]
            assert r["paid_ahead_paise"] == student["paid_ahead_paise"]
            line = next(
                (m for m in details[r["student_id"]]["months"] if m["month"] == month), None
            )
            if line is None:  # a later month the profile doesn't list yet: nothing paid
                assert month > body["current_month"], month
                assert r["paid_paise"] == r["covered_by_credit_paise"] == 0
                assert r["short_paise"] == r["fee_paise"]
                continue
            assert r["fee_paise"] == line["expected_paise"]
            assert r["paid_paise"] == line["paid_paise"]
            assert r["paid_direct_paise"] == line["paid_direct_paise"]
            assert r["covered_by_credit_paise"] == line["covered_by_credit_paise"]
            assert r["credit_sources"] == line["credit_sources"]
            assert r["extra_sent"] == line["extra_sent"]
            assert r["extra_unused_paise"] == line["extra_unused_paise"]
            assert r["short_paise"] == line["remaining_paise"]
            assert (
                r["paid_paise"]
                == r["paid_direct_paise"] + r["extra_sent_paise"] + r["extra_unused_paise"]
            )
            assert (
                r["paid_direct_paise"] + r["covered_by_credit_paise"] + r["short_paise"]
                == r["fee_paise"]
            )

        # Exactly the students relevant to M: enrolled in M, money logged for or paying M, or
        # still owing an earlier month.
        assert {r["student_id"] for r in rows} == {
            i for i in students if _relevant(students[i], details[i], month, i in backlog)
        }, month


# --------------------------------------------------------------------------- the real case


def test_month_paid_twice_while_the_one_before_is_unpaid(api: TestClient) -> None:
    """The owner's case: September paid twice, August not at all. Here, June (now) paid
    twice and May unpaid: the second June payment pays May."""
    s = make_student(api, name="Ananya Rao", joined_month="2026-05", phone="98765 43210")
    pay(api, s["id"], "2026-06", 150000, paid_on="2026-06-05")
    second = pay(api, s["id"], "2026-06", 150000, paid_on="2026-06-10")

    june = rows_by_name(report(api, "2026-06"))["Ananya Rao"]
    assert june["status"] == "paid"
    assert june["fee_paise"] == 150000
    assert june["paid_paise"] == 300000
    assert june["paid_direct_paise"] == 150000
    assert june["extra_sent_paise"] == 150000
    assert june["extra_sent"] == [{"to_month": "2026-05", "amount_paise": 150000}]
    assert june["short_paise"] == 0
    assert june["owed_before_paise"] == 0
    assert june["owed_now_paise"] == 0

    may = rows_by_name(report(api, "2026-05"))["Ananya Rao"]
    assert may["status"] == "paid_with_credit"
    assert may["paid_paise"] == 0
    assert may["covered_by_credit_paise"] == 150000
    assert may["credit_sources"] == [
        {
            "payment_id": second["id"],
            "paid_on": "2026-06-10",
            "for_month": "2026-06",
            "amount_paise": 150000,
        }
    ]
    assert may["short_paise"] == 0
    assert_reconciles(api)


# --------------------------------------------------------------------------- statuses


def _scenario(api: TestClient) -> dict[str, int]:
    """A bit of everything, as of June 2026."""
    ids: dict[str, int] = {}

    def student(name: str, **fields: Any) -> int:
        ids[name] = make_student(api, name=name, **fields)["id"]
        return ids[name]

    # Paid, Partial, Unpaid in June; an earlier month owed.
    pay(api, student("Kabir Mehta", joined_month="2026-04"), "2026-06")
    pay(api, ids["Kabir Mehta"], "2026-04")  # May owed
    pay(api, student("Meera Iyer", joined_month="2026-06"), "2026-06", 50000)
    student("Arjun Menon", joined_month="2026-05")  # May and June unpaid
    # A free place: never owed, never counted.
    student("Zoya Khan", monthly_fee_paise=0, joined_month="2026-03")
    # Left after April, April unpaid, and ₹1,000 logged for May after leaving: all extra,
    # so it pays ₹1,000 of April. ₹500 of April is still owed.
    rohan = student("Rohan Das", joined_month="2026-03", left_month="2026-04")
    pay(api, rohan, "2026-03")
    pay(api, rohan, "2026-05", 100000, paid_on="2026-05-20")
    # Left after March, nothing owed and no money since: not on later reports.
    pay(api, student("Isha Nair", joined_month="2026-03", left_month="2026-03"), "2026-03")
    # Left in March, owes March: on every later report (the dashboard's backlog).
    student("Dev Patel", joined_month="2026-03", left_month="2026-03")
    # Away: left after February, back from May, so March and April have no fee.
    tara = student("Tara Singh", joined_month="2026-01", left_month="2026-02")
    for m in ("2026-01", "2026-02"):
        pay(api, tara, m)
    back = api.post(f"/api/students/{tara}/return", json={"from_month": "2026-05"})
    assert back.status_code == 200, back.text
    pay(api, tara, "2026-05")
    pay(api, tara, "2026-06")
    # Paid ahead for July; August not yet.
    sana = student("Sana Gupta", joined_month="2026-06")
    pay(api, sana, "2026-06")
    pay(api, sana, "2026-07", paid_on="2026-06-15")
    # Joins in August.
    student("Neel Bose", joined_month="2026-08")
    return ids


def test_statuses_and_who_is_on_the_report(api: TestClient) -> None:
    _scenario(api)
    june = report(api, "2026-06")
    assert june["month"] == "2026-06"
    assert june["current_month"] == "2026-06"
    assert june["today"] == FROZEN_TODAY.isoformat()
    by_name = rows_by_name(june)
    assert {n: r["status"] for n, r in by_name.items()} == {
        "Arjun Menon": "unpaid",
        "Meera Iyer": "partial",
        "Kabir Mehta": "paid",
        "Sana Gupta": "paid",
        "Tara Singh": "paid",
        "Zoya Khan": "no_fee",
        "Dev Patel": "left",  # left after March, still owes March
        "Rohan Das": "left",  # left after April, still owes ₹500 of April
    }
    # Unpaid first, then by status, then by name.
    assert [r["student_name"] for r in june["rows"]] == [
        "Arjun Menon",
        "Meera Iyer",
        "Kabir Mehta",
        "Sana Gupta",
        "Tara Singh",
        "Zoya Khan",
        "Dev Patel",
        "Rohan Das",
    ]
    arjun = by_name["Arjun Menon"]
    assert (arjun["short_paise"], arjun["owed_before_paise"], arjun["owed_now_paise"]) == (
        150000,
        150000,
        300000,
    )
    assert arjun["owed_before_months"] == ["2026-05"]
    assert by_name["Kabir Mehta"]["owed_before_months"] == ["2026-05"]
    assert by_name["Sana Gupta"]["paid_ahead_paise"] == 150000
    dev = by_name["Dev Patel"]
    assert (dev["is_enrolled"], dev["fee_paise"], dev["owed_before_paise"]) == (False, 0, 150000)
    zoya = by_name["Zoya Khan"]
    assert (zoya["is_enrolled"], zoya["fee_paise"]) == (True, 0)
    assert by_name["Rohan Das"]["owed_before_paise"] == 50000
    assert june["totals"]["active_student_count"] == 5  # Zoya's free place isn't counted
    assert june["totals"]["not_fully_paid_count"] == 2

    # May: Rohan left after April, but money logged for May pays his April.
    may = rows_by_name(report(api, "2026-05"))
    rohan = may["Rohan Das"]
    assert rohan["status"] == "left"
    assert (rohan["paid_paise"], rohan["fee_paise"], rohan["short_paise"]) == (100000, 0, 0)
    assert rohan["extra_sent"] == [{"to_month": "2026-04", "amount_paise": 100000}]
    assert "Isha Nair" not in may  # left, nothing owed, no money for May
    assert may["Tara Singh"]["status"] == "paid"
    april = rows_by_name(report(api, "2026-04"))
    assert april["Tara Singh"]["status"] == "no_fee"  # a month away
    assert april["Rohan Das"]["status"] == "partial"  # ₹1,000 of ₹1,500, from May's money
    assert april["Rohan Das"]["covered_by_credit_paise"] == 100000

    # July isn't due yet: Sana paid ahead, the rest not due yet. Neel hasn't joined.
    july = rows_by_name(report(api, "2026-07"))
    assert july["Sana Gupta"]["status"] == "paid"
    assert july["Arjun Menon"]["status"] == "not_due_yet"
    assert july["Arjun Menon"]["short_paise"] == 150000
    assert "Neel Bose" not in july
    august = rows_by_name(report(api, "2026-08"))
    assert august["Neel Bose"]["status"] == "not_due_yet"
    assert august["Neel Bose"]["owed_before_paise"] == 0  # nothing before August is due for him

    assert_reconciles(api)


def test_paid_ahead_with_extra_money(api: TestClient) -> None:
    """Extra money that pays a later month makes it Paid with credit there."""
    s = make_student(api, name="Kabir Mehta", joined_month="2026-06")["id"]
    pay(api, s, "2026-06", 300000)
    july = rows_by_name(report(api, "2026-07"))["Kabir Mehta"]
    assert july["status"] == "paid_with_credit"
    assert july["covered_by_credit_paise"] == 150000
    assert_reconciles(api)


def test_money_nobody_needed_is_kept_as_credit(api: TestClient) -> None:
    s = make_student(api, name="Kabir Mehta", joined_month="2026-06", left_month="2026-06")["id"]
    pay(api, s, "2026-06", 200000)
    june = rows_by_name(report(api, "2026-06"))["Kabir Mehta"]
    assert june["status"] == "paid"
    assert june["extra_unused_paise"] == 50000
    assert june["credit_paise"] == 50000
    assert_reconciles(api)


def test_empty_report_and_default_month(api: TestClient) -> None:
    body = report(api)
    assert body["month"] == "2026-06"
    assert body["rows"] == []
    assert body["totals"]["student_count"] == 0
    assert body["totals"]["fee_paise"] == 0


@pytest.mark.parametrize("month", ["2026-13", "June", "2026-6", "1999-12"])
def test_bad_month_is_a_422(api: TestClient, month: str) -> None:
    for path in ("/api/report", "/api/report.xlsx"):
        response = api.get(path, params={"month": month})
        assert response.status_code == 422, (path, month)
        assert response.json()["detail"][0]["loc"] == ["query", "month"]


def test_demo_data_reconciles(api: TestClient) -> None:
    with session_factory()() as session:
        seed(session, FROZEN_TODAY)
    assert_reconciles(api)


# --------------------------------------------------------------------------- Excel


def _sheet(content: bytes) -> Any:
    return load_workbook(io.BytesIO(content)).worksheets[0]


def test_excel_has_the_screen_rows_and_totals(api: TestClient) -> None:
    _scenario(api)
    make_student(api, name="=HYPERLINK(1)", joined_month="2026-06", batch_label="Tue 5pm")
    response = api.get("/api/report.xlsx", params={"month": "2026-06"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "scrappy-records-report-2026-06.xlsx" in response.headers["content-disposition"]
    ws = _sheet(response.content)
    body = report(api, "2026-06")

    assert ws["A1"].value == "Scrappy Records — Fees report, June 2026"
    assert ws["A1"].font.bold
    assert ws["A2"].value == "As of 15 Jun 2026"
    headings = [c.value for c in ws[4]]
    assert headings == [
        "Student",
        "Class/batch",
        "Phone",
        "Fee ₹",
        "Paid for this month ₹ (logged)",
        "Covered by credit ₹",
        "Credit came from",
        "Extra sent elsewhere ₹",
        "Extra went to",
        "Short ₹",
        "Status",
        "Owed from earlier months ₹",
        "Earlier months owed",
        "Total owed now ₹",
        "Credit ₹",
        "Paid ahead ₹",
    ]
    assert all(c.font.bold for c in ws[4])
    assert ws.freeze_panes == "B5"
    assert ws.page_setup.orientation == "landscape"

    data = [[c.value for c in row] for row in ws.iter_rows(min_row=5)]
    rows, total = data[:-1], data[-1]
    assert [r[0] for r in rows] == [r["student_name"] for r in body["rows"]]
    by_name = {r[0]: r for r in rows}
    arjun = by_name["Arjun Menon"]
    assert arjun[3:6] == [1500, 0, 0]
    assert arjun[9:14] == [1500, "Unpaid", 1500, "May 2026", 3000]
    assert by_name["Zoya Khan"][10] == "No fee"
    assert by_name["Dev Patel"][10] == "Left"
    injected = by_name["=HYPERLINK(1)"]
    assert injected[1] == "Tue 5pm"
    cell = next(c for c in ws["A"] if c.value == "=HYPERLINK(1)")
    assert cell.data_type == "s"  # text, never a formula

    t = body["totals"]
    assert total[0] == f"Total ({t['student_count']} students)"
    assert total[3] == t["fee_paise"] // 100
    assert total[4] == t["paid_paise"] // 100
    assert total[9] == t["short_paise"] // 100
    assert total[10] == f"{t['not_fully_paid_count']} of {t['active_student_count']} not fully paid"
    assert total[13] == t["owed_now_paise"] // 100
    total_row = ws.max_row
    assert all(c.font.bold for c in ws[total_row])
    # Money is a number with a ₹ format, so it adds up in Excel.
    fee_cell = ws.cell(row=5, column=4)
    assert isinstance(fee_cell.value, int)
    assert "₹" in fee_cell.number_format
    # The money columns add up to the totals row.
    for col in (4, 5, 6, 8, 10, 12, 14, 15, 16):
        assert sum(r[col - 1] or 0 for r in rows) == total[col - 1], col


def test_excel_credit_words(api: TestClient) -> None:
    s = make_student(api, name="Ananya Rao", joined_month="2026-05")
    pay(api, s["id"], "2026-06", 150000, paid_on="2026-06-05")
    pay(api, s["id"], "2026-06", 170050, paid_on="2026-06-10")
    ws = _sheet(api.get("/api/report.xlsx", params={"month": "2026-05"}).content)
    may = [c.value for c in ws[5]]
    assert may[5] == 1500
    assert may[6] == "₹1,500 from the 10 Jun 2026 payment (for Jun 2026)"
    assert may[10] == "Paid with credit"
    ws = _sheet(api.get("/api/report.xlsx", params={"month": "2026-06"}).content)
    june = ws[5]
    assert june[4].value == 3200.5
    assert june[4].number_format.startswith('"₹"#,##0.00')
    assert june[8].value == "₹1,500 → May 2026; ₹200.50 → Jul 2026"


def test_excel_for_a_month_not_due_yet_says_so(api: TestClient) -> None:
    ws = _sheet(api.get("/api/report.xlsx", params={"month": "2026-08"}).content)
    assert ws["A1"].value == "Scrappy Records — Fees report, August 2026"
    assert "August 2026 isn't due yet" in ws["A2"].value
    assert ws["A5"].value == "Total (0 students)"


def test_rupees_in_indian_grouping() -> None:
    from app.services.report_xlsx import rupees

    assert rupees(0) == "₹0"
    assert rupees(150000) == "₹1,500"
    assert rupees(15000000) == "₹1,50,000"
    assert rupees(1234567890) == "₹1,23,45,678.90"
    assert rupees(150050) == "₹1,500.50"
