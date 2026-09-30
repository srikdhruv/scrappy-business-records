"""The monthly report and unassigned payments: they count for no student, so the report (on
screen and in its Excel file) says how much for its month is still waiting."""

from __future__ import annotations

import io

from fastapi.testclient import TestClient
from helpers import make_student
from openpyxl import load_workbook
from test_unassigned import upload_strays


def _notes(content: bytes) -> list[str]:
    ws = load_workbook(io.BytesIO(content)).active
    assert ws is not None
    return [c.value for c in ws["A"] if isinstance(c.value, str) and "not yet matched" in c.value]


def test_the_monthly_report_says_what_is_waiting(api: TestClient) -> None:
    make_student(api, name="Kabir Mehta", joined_month="2026-06")
    upload_strays(
        api,
        ["Nobody", None, 1500, "2026-06-02", "Jun 2026", "UPI", None],
        ["Someone", None, 3000, "2026-06-03", "Jun 2026", "Cash", None],
        ["Else", None, 900, "2026-05-03", "May 2026", "Cash", None],
    )
    june = api.get("/api/report", params={"month": "2026-06"}).json()
    assert (june["unassigned_count"], june["unassigned_paise"]) == (2, 450000)
    assert june["totals"]["paid_paise"] == 0  # not counted for anyone
    july = api.get("/api/report", params={"month": "2026-07"}).json()
    assert (july["unassigned_count"], july["unassigned_paise"]) == (0, 0)

    june_file = api.get("/api/report.xlsx", params={"month": "2026-06"}).content
    assert _notes(june_file) == [
        "Also ₹4,500 of payments not yet matched to a student (2 payments for June 2026 from an "
        "upload): not counted above. Give them to a student in Payments → Unassigned payments."
    ]
    assert _notes(api.get("/api/report.xlsx", params={"month": "2026-07"}).content) == []
