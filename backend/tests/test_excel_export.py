"""Excel downloads, and the round trip: Download everything, then upload it into an empty app,
gives back exactly the same records (app/services/exports.py)."""

from __future__ import annotations

import datetime as dt
import io
import json
from collections import Counter
from typing import Any

from conftest import FROZEN_TODAY
from fastapi.testclient import TestClient
from helpers import make_student, pay
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from test_excel_import import commit, preview, xlsx

from app.db import session_factory
from app.seed import seed

Json = dict[str, Any]


def sheet(content: bytes, title: str | None = None) -> Worksheet:
    book = load_workbook(io.BytesIO(content))
    return book[title] if title else book.worksheets[0]


def values(ws: Worksheet) -> list[list[Any]]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


# --------------------------------------------------------------------------- the round trip


_IDS = ("id", "student_id", "payment_id", "batch_id", "created_at", "updated_at")


def _strip(item: Any) -> Any:
    """`item` without database ids or timestamps, at any depth (a credit source names the
    payment it came from by id)."""
    if isinstance(item, dict):
        return {k: _strip(v) for k, v in item.items() if k not in _IDS}
    if isinstance(item, list):
        return [_strip(v) for v in item]
    return item


def _records(api: TestClient) -> Json:
    """Everything the app knows, without ids or timestamps, so two apps can be compared: every
    student in full (their ledger month by month), every payment with who it belongs to, every
    unassigned payment, and the Dashboard for every month from the first payment to a year
    ahead. Students with the same name stay apart (as a multiset of whole records)."""
    students = api.get("/api/students", params={"status": "all"}).json()
    details: dict[int, Json] = {}
    for s in students:
        d = _strip(api.get(f"/api/students/{s['id']}").json())
        d["fee_history"] = [_strip(f) for f in d["fee_history"]]
        if d["next_fee_change"]:
            d["next_fee_change"] = _strip(d["next_fee_change"])
        details[s["id"]] = d
    who = {
        sid: json.dumps({k: d[k] for k in ("name", "phone", "joined_month", "batch_label",
                                           "notes", "guardian_name")}, sort_keys=True)
        for sid, d in details.items()
    }  # fmt: skip
    payments = Counter(
        (who[p["student_id"]], p["amount_paise"], p["paid_on"], p["for_month"], p["method"],
         p["note"])
        for p in api.get("/api/payments").json()
    )  # fmt: skip
    unassigned = Counter(
        tuple(u[k] for k in ("student_text", "phone", "amount_paise", "paid_on", "for_month",
                             "method", "note", "source"))
        for u in api.get("/api/unassigned-payments").json()
    )  # fmt: skip
    dashboards = {}
    month = "2024-01"
    while month <= "2027-06":
        board = api.get("/api/dashboard", params={"month": month}).json()
        dashboards[month] = {
            "summary": board["summary"],
            "lists": [
                json.dumps(_strip(item), sort_keys=True)
                for part in ("yet_to_pay", "backlog", "overpaid")
                for item in board[part]
            ],
        }
        year, mon = int(month[:4]), int(month[5:])
        month = f"{year + mon // 12}-{mon % 12 + 1:02d}"
    batches = [_strip(b) for b in api.get("/api/batches").json()]
    return {
        "students": sorted(json.dumps(d, sort_keys=True) for d in details.values()),
        "payments": payments,
        "unassigned": unassigned,
        "dashboards": dashboards,
        "batches": batches,
    }


def _rich_records(api: TestClient) -> None:
    """The demo data (five batches, everyone in one), plus a return after leaving (months
    away), a month off, a planned fee, notes with odd characters, an unassigned payment, a
    batch with every detail and odd characters, an empty batch, a student in no batch, and an
    old class label."""
    with session_factory()() as session:
        seed(session, FROZEN_TODAY)
    odd = api.post(
        "/api/batches",
        json={
            "name": "Café Seniors \N{EN DASH} रविवार",
            "location": '=Studio "B"',
            "days": ["sat", "sun"],
            "start_time": "09:30",
            "end_time": "11:00",
            "default_fee_paise": 149950,
            "notes": "Two lines\nhere",
        },
    ).json()
    api.post("/api/batches", json={"name": "Nobody yet"})
    make_student(api, name="Tara", batch_id=odd["id"], batch_label="Old: Sat 9am")
    make_student(api, name="Kiran Bose", batch_label="Thu 7pm")  # no batch, a label
    back = make_student(
        api,
        name="Émile O'Brien",
        phone="+91 98765-43210",
        joined_month="2025-06",
        left_month="2025-10",
        notes="=SUM(A1:A3) is text, not a formula\nSecond line",
    )
    assert api.post(f"/api/students/{back['id']}/return", json={"from_month": "2026-02"}).is_success
    api.patch(
        f"/api/students/{back['id']}",
        json={"monthly_fee_paise": 0, "fee_effective_month": "2026-07"},
    )
    api.patch(
        f"/api/students/{back['id']}",
        json={"monthly_fee_paise": 175050, "fee_effective_month": "2026-08"},
    )
    pay(api, back["id"], "2026-03", 150000, note="half & half ₹")
    pay(api, back["id"], "2025-12", 20000)  # paid while away: credit
    stray = xlsx(
        ("Payments", [["Student", "Amount", "Paid on", "For month", "Method"],
                      ["Someone Unknown", "₹999.50", "3 Jun 2026", "Jun 2026", "Cash"]]),
    )  # fmt: skip
    shown = preview(api, stray, "bank statement.xlsx")
    assert commit(api, shown)["unassigned_added"] == 1


def _wipe(api: TestClient) -> None:
    for s in api.get("/api/students", params={"status": "all"}).json():
        assert api.delete(f"/api/students/{s['id']}").status_code == 204
    for u in api.get("/api/unassigned-payments").json():
        assert api.delete(f"/api/unassigned-payments/{u['id']}").status_code == 204
    for b in api.get("/api/batches").json():
        assert api.delete(f"/api/batches/{b['id']}").status_code == 204


def test_download_everything_restores_everything_into_an_empty_app(api: TestClient) -> None:
    _rich_records(api)
    before = _records(api)
    assert len(before["students"]) == 28 and sum(before["unassigned"].values()) == 1
    assert len(before["batches"]) == 7
    response = api.get("/api/export/everything.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        'filename="scrappy-records-everything-2026-06-15.xlsx"'
        in (response.headers["content-disposition"])
    )
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == [
        "Students",
        "Batches",
        "Fee history",
        "Payments",
        "Unassigned payments",
    ]

    _wipe(api)
    assert api.get("/api/students", params={"status": "all"}).json() == []

    shown = preview(api, response.content, "everything.xlsx")
    assert shown["ignored_sheets"] == []
    assert {s["status"] for s in shown["students"]} == {"new"}
    assert {p["status"] for p in shown["payments"]} == {"ready", "unassigned"}
    assert {b["status"] for b in shown["batches"]} == {"new"}
    assert len(shown["batches"]) == 7
    result = commit(api, shown)
    assert result["students_added"] == 28
    assert result["batches_added"] == 7
    assert _records(api) == before

    # Uploading it again adds nothing: everything is already here.
    again = preview(api, response.content)
    assert {s["status"] for s in again["students"]} == {"exists"}
    assert {p["status"] for p in again["payments"]} == {"duplicate"}
    assert {b["status"] for b in again["batches"]} == {"exists"}
    assert commit(api, again)["backup_file"] is None
    assert _records(api) == before


def _look_alikes(api: TestClient) -> None:
    """People an upload could easily mix up: two "Priya S" with no phone, brothers and sisters
    sharing a parent's phone, students with no phone at all, a hyphenated name, and two
    identical payments on one day (two instalments)."""
    first = make_student(api, name="Priya S", joined_month="2026-01", batch_label="Mon")
    second = make_student(api, name="Priya S", joined_month="2026-03", batch_label="Tue")
    for student, month in ((first, "2026-01"), (first, "2026-02"), (second, "2026-03")):
        pay(api, student["id"], month)
    sisters = [
        make_student(api, name=name, phone="98765 43210", guardian_name="Lakshmi Iyer",
                     joined_month="2026-02", monthly_fee_paise=120000)
        for name in ("Tara Iyer", "Meera Iyer", "Anu Iyer")
    ]  # fmt: skip
    for sister in sisters:
        pay(api, sister["id"], "2026-02", 120000)
        pay(api, sister["id"], "2026-03", 60000, paid_on="2026-03-04")
        pay(api, sister["id"], "2026-03", 60000, paid_on="2026-03-04")  # the same, twice
    lone = make_student(api, name="Mary-Jane Dsouza", joined_month="2026-04")
    pay(api, lone["id"], "2026-04", 150000, note="=cash")


def test_download_everything_keeps_look_alikes_apart(api: TestClient) -> None:
    """Restoring into an empty app never merges or drops anyone, even people with the same
    name, or the same phone, and genuine repeated payments all come back."""
    _look_alikes(api)
    before = _records(api)
    content = api.get("/api/export/everything.xlsx").content
    _wipe(api)

    shown = preview(api, content)
    assert [s["status"] for s in shown["students"]] == ["new"] * 6
    assert {p["status"] for p in shown["payments"]} == {"ready"}
    result = commit(api, shown)
    assert (result["students_added"], result["payments_added"]) == (6, 13)
    assert _records(api) == before

    # Back into the same app: every row is recognised as already here, and nothing is added.
    again = preview(api, content)
    assert [s["status"] for s in again["students"]] == ["exists"] * 6
    assert {p["status"] for p in again["payments"]} == {"duplicate"}
    assert commit(api, again)["backup_file"] is None
    assert _records(api) == before


def test_everything_into_an_app_with_some_students_already(api: TestClient) -> None:
    kabir = make_student(api, name="Kabir Mehta", phone="90000 00001")
    pay(api, kabir["id"], "2026-01")
    content = api.get("/api/export/everything.xlsx").content
    make_student(api, name="Diya Nair")  # only in this app
    # Kabir's fee changed here after the download: it isn't overwritten.
    api.patch(f"/api/students/{kabir['id']}", json={"monthly_fee_paise": 200000})
    shown = preview(api, content)
    [row] = shown["students"]
    assert row["status"] == "exists" and row["student_id"] == kabir["id"]
    assert shown["fee_changes"] == 0  # an existing student's fee history is left alone
    assert shown["payments"][0]["status"] == "duplicate"
    commit(api, shown)
    assert api.get(f"/api/students/{kabir['id']}").json()["monthly_fee_paise"] == 200000
    assert len(api.get("/api/payments").json()) == 1


# --------------------------------------------------------------------------- what's shown


def test_students_download_matches_the_tab_and_search(api: TestClient) -> None:
    make_student(api, name="Arjun Menon", phone="90000 00006", batch_label="Tue 5pm")
    make_student(api, name="Émile Dsouza", phone="+91 98765-43210", notes="Loves jazz")
    make_student(api, name="Rohan Desai", joined_month="2025-01", left_month="2025-12")
    ananya = make_student(api, name="Ananya Rao", guardian_name="Lakshmi Rao")
    pay(api, ananya["id"], "2026-01", 100000)

    def names(**params: str) -> list[str]:
        response = api.get("/api/export/students.xlsx", params=params)
        assert response.status_code == 200
        return [r[0] for r in values(sheet(response.content))[1:]]

    assert names() == ["Ananya Rao", "Arjun Menon", "Émile Dsouza"]  # active by default
    assert names(status="left") == ["Rohan Desai"]
    assert names(status="all", q="emile") == ["Émile Dsouza"]
    assert names(status="all", q="menon arjun") == ["Arjun Menon"]  # the page's search
    assert names(status="all", q="9876543210") == ["Émile Dsouza"]
    assert names(status="all", q="lakshmi") == ["Ananya Rao"]
    assert names(status="all", q="5pm tue") == ["Arjun Menon"]

    ws = sheet(api.get("/api/export/students.xlsx", params={"q": "ananya"}).content)
    assert ws.title == "Students"
    assert values(ws) == [
        ["Name", "Phone", "Parent/guardian", "Batch", "Old class label",
         "Monthly fee ₹ (current)", "Joined (month)", "Left (month)", "Status", "Owes ₹",
         "Notes"],
        ["Ananya Rao", None, "Lakshmi Rao", None, None, 1500, dt.datetime(2026, 1, 1), None,
         "Owes", 8000, None],
    ]  # fmt: skip
    assert ws.freeze_panes == "A2"
    assert ws["A1"].font.bold
    assert ws["F2"].number_format == '"₹"#,##0'
    assert ws["G2"].number_format == "mmm yyyy"
    assert ws.column_dimensions["A"].width >= 20


def test_payments_download_matches_the_filters_and_sort(api: TestClient) -> None:
    ananya = make_student(api, name="Ananya Rao", phone="98765 43210")
    kabir = make_student(api, name="Kabir Mehta")
    pay(api, ananya["id"], "2026-01", 150000, paid_on="2026-01-05", method="cash")
    pay(api, ananya["id"], "2026-02", 120050, paid_on="2026-02-03", note="=1+1")
    pay(api, kabir["id"], "2026-02", 90000, paid_on="2026-02-10", method="cash")

    def rows(**params: str) -> list[list[Any]]:
        response = api.get("/api/export/payments.xlsx", params=params)
        assert response.status_code == 200, response.text
        return values(sheet(response.content))

    head, *everything = rows()
    assert head == ["Student", "Phone", "Amount ₹", "Paid on", "For month", "Method", "Note"]
    assert [r[3] for r in everything] == [
        dt.datetime(2026, 2, 10),
        dt.datetime(2026, 2, 3),
        dt.datetime(2026, 1, 5),
    ]  # newest first, like the page
    assert everything[1] == [
        "Ananya Rao", "98765 43210", 1200.5, dt.datetime(2026, 2, 3), dt.datetime(2026, 2, 1),
        "UPI", "=1+1",
    ]  # fmt: skip
    assert [r[0] for r in rows(method="cash")[1:]] == ["Kabir Mehta", "Ananya Rao"]
    assert [r[2] for r in rows(sort="amount", order="asc")[1:]] == [900, 1200.5, 1500]
    assert [r[0] for r in rows(month="2026-02", q="ananya")[1:]] == ["Ananya Rao"]
    assert [r[2] for r in rows(student_id=str(kabir["id"]))[1:]] == [900]

    ws = sheet(api.get("/api/export/payments.xlsx").content)
    assert ws["G3"].data_type == "s"  # "=1+1" stays a note, never a formula
    assert ws["C3"].number_format == '"₹"#,##0.00'
    assert ws["D2"].number_format == "d mmm yyyy"


def test_downloads_reject_bad_filters_plainly(api: TestClient) -> None:
    for path, params in (
        ("/api/export/payments.xlsx", {"month": "2026-13"}),
        ("/api/export/payments.xlsx", {"method": "cheque"}),
        ("/api/export/payments.xlsx", {"sort": "colour"}),
        ("/api/export/students.xlsx", {"status": "gone"}),
        ("/api/import/template.xlsx", {"kind": "everything"}),
    ):
        assert api.get(path, params=params).status_code == 422, (path, params)
