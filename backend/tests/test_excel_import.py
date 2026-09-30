"""Uploading Excel files: preview first (nothing saved), then add (checked again, one
transaction, a backup first). See app/services/imports.py."""

from __future__ import annotations

import base64
import datetime as dt
import io
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from helpers import make_student, pay
from openpyxl import Workbook, load_workbook

from app import config

Json = dict[str, Any]


def xlsx(*sheets: tuple[str, list[list[Any]]]) -> bytes:
    """A workbook with these sheets: (title, rows), the first row being the headings."""
    book = Workbook()
    book.remove(book.active)  # type: ignore[arg-type]
    for title, rows in sheets:
        ws = book.create_sheet(title)
        for row in rows:
            ws.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def preview(api: TestClient, data: bytes, filename: str = "upload.xlsx") -> Json:
    response = api.post("/api/import/preview", params={"filename": filename}, content=data)
    assert response.status_code == 200, response.text
    shown = response.json()
    shown["_file"] = data  # Add sends the file again (not part of the API's answer)
    return shown


def preview_error(api: TestClient, data: bytes) -> str:
    response = api.post("/api/import/preview", content=data)
    assert response.status_code == 422, response.text
    [item] = response.json()["detail"]
    return item["msg"]


def commit_body(
    shown: Json,
    students: dict[int, bool] | None = None,
    payments: dict[int, dict[str, Any]] | None = None,
) -> Json:
    """What Add sends: the file again, and choices by row number (students: add; payments:
    choice and student_id)."""
    sheets = {p["row"]: p["sheet"] for p in shown["payments"]}
    return {
        "file": base64.b64encode(shown["_file"]).decode(),
        "filename": shown["filename"],
        "students": [{"row": row, "add": add} for row, add in (students or {}).items()],
        "payments": [
            {"sheet": sheets.get(row, "Payments"), "row": row, **choice}
            for row, choice in (payments or {}).items()
        ],
    }


def commit(
    api: TestClient,
    shown: Json,
    students: dict[int, bool] | None = None,
    payments: dict[int, dict[str, Any]] | None = None,
) -> Json:
    """Add a preview's file, with choices by row number."""
    response = api.post("/api/import/commit", json=commit_body(shown, students, payments))
    assert response.status_code == 200, response.text
    return response.json()


def by_row(rows: list[Json]) -> dict[int, Json]:
    return {r["row"]: r for r in rows}


STUDENT_HEAD = ["Name", "Phone", "Monthly fee", "Joined"]
PAYMENT_HEAD = ["Student", "Amount", "Paid on", "For month", "Method"]


# --------------------------------------------------------------------------- students


def test_student_rows_are_classified(api: TestClient) -> None:
    make_student(api, name="Ananya Rao", phone="98765 43210")
    make_student(api, name="Kabir Mehta")  # no phone
    make_student(api, name="Meera Iyer", phone="90000 00001")
    make_student(api, name="Tanvi Shetty", phone="90000 00002")
    data = xlsx(
        (
            "Sheet1",
            [
                STUDENT_HEAD,
                ["rao ananya", "+91 98765-43210", 1500, "Jan 2026"],  # 2 exists: same name+phone
                ["Kabir Mehta", None, 1500, "Jan 2026"],  # 3 exists: same name, no phones
                ["Meera Iyer", "90000 00009", 1500, "Jan 2026"],  # 4 similar: other phone
                ["Rohan Desai", "90000 00002", 1500, "Jan 2026"],  # 5 similar: Tanvi's phone
                ["Ishaan Kapoor", "90000 00003", "₹1,800", "Feb 2026"],  # 6 new
                ["Ishaan Kapoor", "90000 00003", 1800, "Feb 2026"],  # 7 repeat of row 6
                ["Diya Nair", None, None, "Feb 2026"],  # 8 problem: no fee
                ["", None, 1500, None],  # 9 problem: no name
                ["Émile Dsouza", None, 1500, "13/2026"],  # 10 problem: bad month
                ["Zara Khan", None, 1500, "Jan 2030"],  # 11 problem: too far ahead
                ["Aarav Joshi", None, "abc", None],  # 12 problem: fee isn't money
            ],
        )
    )
    shown = preview(api, data)
    rows = by_row(shown["students"])
    assert {r: rows[r]["status"] for r in rows} == {
        2: "exists",
        3: "exists",
        4: "similar",
        5: "similar",
        6: "new",
        7: "exists",
        8: "problem",
        9: "problem",
        10: "problem",
        11: "problem",
        12: "problem",
    }
    assert rows[2]["reason"] == "Already here: Ananya Rao (98765 43210)"
    assert rows[4]["reason"] == ("Same name as Meera Iyer (90000 00001), but a different phone")
    assert rows[5]["reason"] == "Same phone as Tanvi Shetty (90000 00002), but a different name"
    assert rows[7]["reason"] == "Same as row 6 in this file"
    assert rows[8]["reason"] == "Monthly fee is missing"
    assert rows[9]["reason"] == "Name is missing"
    assert rows[10]["reason"].startswith("Joined month “13/2026” isn't a month")
    assert rows[11]["reason"] == "Joined month can't be later than June 2028 (two years from now)"
    assert rows[12]["reason"].startswith("Monthly fee “abc” isn't an amount")
    assert rows[6]["monthly_fee_paise"] == 180000
    # Nothing was saved by the preview.
    assert len(api.get("/api/students", params={"status": "all"}).json()) == 4

    result = commit(api, shown, students={5: True})  # add Rohan anyway; Meera skipped
    assert result["students_added"] == 2
    names = [s["name"] for s in api.get("/api/students", params={"status": "all"}).json()]
    assert names == [
        "Ananya Rao",
        "Ishaan Kapoor",
        "Kabir Mehta",
        "Meera Iyer",
        "Rohan Desai",
        "Tanvi Shetty",
    ]
    # "Already here" changes nothing about the student who is here.
    ananya = api.get("/api/students", params={"q": "ananya"}).json()[0]
    assert ananya["phone"] == "98765 43210" and ananya["joined_month"] == "2026-01"


def test_blank_joined_month_is_this_month_and_optional_columns_come_through(
    api: TestClient,
) -> None:
    data = xlsx(
        (
            "Students",
            [
                ["NAME", "Mobile", "Parent", "Batch", "Fee", "Notes", "Left"],
                ["Kabir Mehta", 9876543210, "Lakshmi Mehta", "Sat 10am", "1500/-", "Keen", None],
            ],
        )
    )
    shown = preview(api, data)
    [row] = shown["students"]
    assert row["status"] == "new"
    assert (row["name"], row["phone"], row["monthly_fee_paise"], row["joined_month"]) == (
        "Kabir Mehta",
        "9876543210",
        150000,
        "2026-06",
    )
    commit(api, shown)
    [student] = api.get("/api/students").json()
    assert {k: student[k] for k in ("name", "phone", "guardian_name", "batch_label", "notes",
                                    "joined_month", "left_month", "monthly_fee_paise")} == {
        "name": "Kabir Mehta",
        "phone": "9876543210",
        "guardian_name": "Lakshmi Mehta",
        "batch_label": "Sat 10am",
        "notes": "Keen",
        "joined_month": "2026-06",
        "left_month": None,
        "monthly_fee_paise": 150000,
    }  # fmt: skip


# --------------------------------------------------------------------------- payments


def test_payment_rows_are_matched_and_classified(api: TestClient) -> None:
    ananya = make_student(api, name="Ananya Rao", phone="98765 43210")
    make_student(api, name="Kabir Mehta", phone="90000 00001")
    make_student(api, name="Kabir Mehta", phone="90000 00002")
    make_student(api, name="Meera Iyer")
    pay(api, ananya["id"], "2026-03")  # ₹1,500 paid 5 Mar 2026 for March
    data = xlsx(
        (
            "Payments",
            [
                [*PAYMENT_HEAD, "Phone"],
                ["ananya rao", 1500, "05/04/2026", "Apr 2026", "UPI", None],  # 2 ready
                ["9876543210", 1500, "5 May 2026", "2026-05", "gpay", None],  # 3 ready by phone
                ["Kabir Mehta", 1500, "2026-05-02", "May 2026", "Cash", None],  # 4 ambiguous
                ["Kabir Mehta", 1500, "2026-05-02", "May 2026", "Cash", "90000 00002"],  # 5 ready
                ["Anaya R", 1500, "2026-05-02", "May 2026", "Cash", None],  # 6 no match
                ["Ananya Rao", 1500, "05/03/2026", "Mar 2026", "UPI", None],  # 7 already logged
                ["ANANYA RAO", 1500, "05/04/2026", "Apr 2026", "UPI", None],  # 8 same as row 2
                ["Meera Iyer", 0, "05/04/2026", "Apr 2026", None, None],  # 9 problem: ₹0
                ["Meera Iyer", 1500, "05/04/2030", "Apr 2026", None, None],  # 10 future
                ["Meera Iyer", 1500, "yesterday", "Apr 2026", None, None],  # 11 bad date
                ["Meera Iyer", "₹10,00,001", "05/04/2026", "Apr 2026", None, None],  # 12 cap
                ["Meera Iyer", 1500, "05/04/2026", None, "bank", None],  # 13 month from date
            ],
        )
    )
    shown = preview(api, data, "april.xlsx")
    rows = by_row(shown["payments"])
    status = {r: rows[r]["status"] for r in rows}
    assert status == {
        2: "ready",
        3: "ready",
        4: "needs_student",
        5: "ready",
        6: "needs_student",
        7: "duplicate",
        8: "duplicate",
        9: "problem",
        10: "problem",
        11: "problem",
        12: "problem",
        13: "ready",
    }
    assert rows[2]["student_id"] == ananya["id"] and rows[3]["student_id"] == ananya["id"]
    assert rows[3]["method"] == "upi" and rows[13]["method"] == "other"
    assert len(rows[4]["candidate_ids"]) == 2
    assert rows[4]["reason"].startswith("More than one student matches “Kabir Mehta”")
    assert rows[6]["reason"].startswith("No student called “Anaya R”")
    assert rows[6]["candidate_ids"] == []  # "Anaya" is no one's word
    assert rows[7]["reason"] == "Already logged: ₹1,500 paid on 5 Mar 2026 for Mar 2026"
    assert rows[8]["reason"] == "Same as row 2 in this file"
    assert rows[9]["reason"] == "Amount must be more than ₹0"
    assert rows[10]["reason"] == "Paid-on date can't be in the future"
    assert rows[11]["reason"].startswith("Paid-on date “yesterday” isn't a date")
    assert rows[12]["reason"] == "Amount is more than ₹10,00,000, the most it can be"
    assert rows[13]["for_month"] == "2026-04"

    kabir_2 = rows[5]["student_id"]
    result = commit(api, shown, payments={6: {"choice": "skip"}})
    assert result == {
        "students_added": 0,
        "fee_changes_added": 0,
        "payments_added": 4,  # rows 2, 3, 5, 13
        "unassigned_added": 1,  # row 4 (ambiguous), kept by default
        "skipped": 7,  # row 6 (skipped), 7 and 8 (duplicates), 9-12 (problems)
        "backup_file": result["backup_file"],
    }
    assert result["backup_file"].startswith("records-pre-import-")
    [waiting] = api.get("/api/unassigned-payments").json()
    assert waiting["student_text"] == "Kabir Mehta"
    assert waiting["source"] == "Upload: april.xlsx"
    assert len(waiting["suggested_student_ids"]) == 2
    kabir_payments = api.get("/api/payments", params={"student_id": kabir_2}).json()
    assert [(p["for_month"], p["method"]) for p in kabir_payments] == [("2026-05", "cash")]


def test_choosing_a_student_for_an_unmatched_payment(api: TestClient) -> None:
    kabir = make_student(api, name="Kabir Mehta")
    data = xlsx(("Payments", [PAYMENT_HEAD, ["K. Mehta", 1500, "2026-05-02", "May 2026", "UPI"]]))
    shown = preview(api, data)
    [row] = shown["payments"]
    assert row["status"] == "needs_student"
    assert row["candidate_ids"] == [kabir["id"]]  # "Mehta" finds him
    commit(api, shown, payments={2: {"choice": "student", "student_id": kabir["id"]}})
    assert len(api.get("/api/payments", params={"student_id": kabir["id"]}).json()) == 1
    assert api.get("/api/unassigned-payments").json() == []


def test_payments_follow_new_students_in_the_same_file(api: TestClient) -> None:
    make_student(api, name="Meera Iyer", phone="90000 00001")
    data = xlsx(
        (
            "Students",
            [
                STUDENT_HEAD,
                ["Diya Nair", None, 1200, "Mar 2026"],
                ["Meera Iyer", "90000 00009", 1500, "Mar 2026"],
            ],
        ),
        (
            "Payments",
            [
                ["Student", "Phone", "Amount", "Paid on", "For month"],
                ["Diya Nair", None, 1200, "2026-03-04", "Mar 2026"],  # the new student
                ["Diya Nair", None, 1200, "2026-03-04", "Mar 2026"],  # duplicate in the file
                ["Meera Iyer", "90000 00009", 1500, "2026-03-04", "Mar 2026"],  # follows row 3
            ],
        ),
    )
    shown = preview(api, data)
    students = by_row(shown["students"])
    assert students[2]["status"] == "new" and students[3]["status"] == "similar"
    payments = by_row(shown["payments"])
    assert payments[2]["status"] == "ready" and payments[2]["student_row"] == 2
    assert payments[3]["status"] == "duplicate"
    assert payments[4]["status"] == "follows_student" and payments[4]["student_row"] == 3

    # Skip the look-alike: its payment is kept as unassigned, not lost.
    result = commit(api, shown)
    assert (result["students_added"], result["payments_added"], result["unassigned_added"]) == (
        1,
        1,
        1,
    )
    diya = api.get("/api/students", params={"q": "diya"}).json()[0]
    assert diya["status"] == "up_to_date" or diya["owed_paise"] >= 0
    assert len(api.get("/api/payments", params={"student_id": diya["id"]}).json()) == 1


def test_adding_the_look_alike_takes_its_payments_with_it(api: TestClient) -> None:
    make_student(api, name="Meera Iyer", phone="90000 00001")
    data = xlsx(
        ("Students", [STUDENT_HEAD, ["Meera Iyer", "90000 00009", 1500, "Mar 2026"]]),
        (
            "Payments",
            [
                ["Student", "Phone", "Amount", "Paid on"],
                ["Meera Iyer", "90000 00009", 1500, "2026-03-04"],
            ],
        ),
    )
    shown = preview(api, data)
    result = commit(api, shown, students={2: True})
    assert (result["students_added"], result["payments_added"]) == (1, 1)
    new = next(s for s in api.get("/api/students").json() if s["phone"] == "90000 00009")
    assert len(api.get("/api/payments", params={"student_id": new["id"]}).json()) == 1


# --------------------------------------------------------------------------- checked again


def test_commit_checks_everything_again(api: TestClient) -> None:
    kabir = make_student(api, name="Kabir Mehta")
    data = xlsx(
        ("Students", [STUDENT_HEAD, ["Diya Nair", None, 1200, "Mar 2026"]]),
        ("Payments", [PAYMENT_HEAD, ["Kabir Mehta", 1500, "2026-05-02", "May 2026", "UPI"]]),
    )
    shown = preview(api, data)
    assert shown["students"][0]["status"] == "new"
    assert shown["payments"][0]["status"] == "ready"
    # Meanwhile: Diya is added by hand, and Kabir's payment is logged.
    make_student(api, name="Diya Nair", joined_month="2026-03", monthly_fee_paise=120000)
    pay(api, kabir["id"], "2026-05", paid_on="2026-05-02")
    result = commit(api, shown)
    assert (result["students_added"], result["payments_added"], result["skipped"]) == (0, 0, 2)
    assert result["backup_file"] is None  # nothing added, so no backup was needed
    assert len(api.get("/api/students").json()) == 2


def test_commit_never_trusts_the_client(api: TestClient) -> None:
    """Add reads the file again: choices can't force in a row that's already here, breaks a
    rule, or isn't in the file at all."""
    make_student(api, name="Kabir Mehta")
    data = xlsx(
        (
            "Students",
            [
                STUDENT_HEAD,
                ["Kabir Mehta", None, 1500, "Jan 2026"],  # 2 already here
                ["Diya Nair", None, 1500, "Jan 2030"],  # 3 too far ahead
            ],
        ),
        (
            "Payments",
            [
                PAYMENT_HEAD,
                ["Nobody", 1, "2031-01-01", "Jan 2026", "UPI"],  # 2 paid in the future
                ["Nobody", 1, "2026-01-01", "Jan 2026", "UPI"],  # 3
            ],
        ),
    )
    shown = preview(api, data)
    body = commit_body(
        shown,
        students={2: True, 3: True, 99: True},
        payments={
            2: {"choice": "student", "student_id": 1},
            3: {"choice": "student", "student_id": 999},  # no such student
            98: {"choice": "add"},
        },
    )
    response = api.post("/api/import/commit", json=body)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["students_added"] == 0
    assert result["payments_added"] == 0
    assert result["unassigned_added"] == 1  # the chosen student doesn't exist: kept, not lost
    assert len(api.get("/api/students").json()) == 1

    # A body that isn't right is a plain 422.
    for bad in (
        {**body, "file": "not base64!"},
        {**body, "file": base64.b64encode(b"Name,Fee").decode()},
        {**body, "students": [{"row": 2, "add": True, "status": "new"}]},
        {**body, "payments": [{"sheet": "Payments", "row": 2, "choice": "student"}]},
        {key: value for key, value in body.items() if key != "file"},
    ):
        assert api.post("/api/import/commit", json=bad).status_code == 422, bad


def test_commit_is_one_transaction(api: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    make_student(api, name="Kabir Mehta")
    data = xlsx(
        ("Students", [STUDENT_HEAD, ["Diya Nair", None, 1200, "Mar 2026"]]),
        ("Payments", [PAYMENT_HEAD, ["Kabir Mehta", 1500, "2026-05-02", "May 2026", "UPI"]]),
    )
    shown = preview(api, data)

    from app.services import imports

    def broken(*_: object, **__: object) -> None:
        raise RuntimeError("disk full")

    # Payments are written after Diya was added; only this patch is undone after the with
    # (not conftest's SCRAPPY_HOME).
    with monkeypatch.context() as patch:
        patch.setattr(imports, "insert", broken)
        with pytest.raises(RuntimeError):
            commit(api, shown)
    names = [s["name"] for s in api.get("/api/students").json()]
    assert names == ["Kabir Mehta"]  # Diya was rolled back
    assert api.get("/api/payments").json() == []
    # And it works once the fault is gone.
    assert commit(api, shown)["students_added"] == 1


def test_a_backup_is_taken_before_adding(api: TestClient) -> None:
    make_student(api, name="Kabir Mehta")
    shown = preview(api, xlsx(("Students", [STUDENT_HEAD, ["Diya Nair", None, 1200, None]])))
    result = commit(api, shown)
    backup = config.backup_dir() / result["backup_file"]
    assert backup.is_file()
    import sqlite3

    with sqlite3.connect(backup) as db:  # the state before the upload
        assert db.execute("SELECT name FROM students").fetchall() == [("Kabir Mehta",)]


def test_no_backup_no_import(api: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import backup

    def fail(_reason: str) -> Path:
        raise OSError("no space")

    monkeypatch.setattr(backup, "backup", fail)
    shown = preview(api, xlsx(("Students", [STUDENT_HEAD, ["Diya Nair", None, 1200, None]])))
    body = commit_body(shown)
    response = api.post("/api/import/commit", json=body)
    assert response.status_code == 422
    assert "Couldn't save a backup first" in response.json()["detail"][0]["msg"]
    assert api.get("/api/students").json() == []


# --------------------------------------------------------------------------- bad files


def test_bad_files_get_plain_messages(api: TestClient) -> None:
    assert preview_error(api, b"") == "The file is empty."
    assert preview_error(api, b"Name,Fee\nKabir,1500\n").startswith("This isn't an Excel file")
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100
    assert "older kind of Excel file (.xls)" in preview_error(api, ole)
    assert preview_error(api, b"PK\x03\x04garbage").startswith("This isn't an Excel file")
    wrong = xlsx(("Sheet1", [["Colour", "Size"], ["red", 3]]))
    assert preview_error(api, wrong).startswith("Couldn't find students or payments")
    empty = xlsx(("Sheet1", []))
    assert preview_error(api, empty).startswith("Couldn't find students or payments")
    huge = b"PK" + b"\0" * (5 * 1024 * 1024)
    assert preview_error(api, huge) == "This file is bigger than 5 MB. Upload a smaller one."


def test_too_many_rows(api: TestClient) -> None:
    rows = [STUDENT_HEAD] + [[f"Student {i}", None, 1000, None] for i in range(5001)]
    assert "more than 5,000 rows" in preview_error(api, xlsx(("Students", rows)))


def test_a_zip_bomb_is_refused(api: TestClient) -> None:
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<x/>")
        z.writestr("xl/worksheets/sheet1.xml", "0" * (90 * 1024 * 1024))
    assert preview_error(api, out.getvalue()).startswith("This file is too big to read")


def test_headings_can_be_below_a_title_and_other_sheets_are_named(api: TestClient) -> None:
    data = xlsx(
        ("Notes", [["Anything"], ["here"]]),
        (
            "Fees Oct",
            [["October fees"], [], PAYMENT_HEAD, ["Nobody Here", 1500, "5/10/2026", None, None]],
        ),
    )
    shown = preview(api, data)
    assert shown["sheets"] == ["Fees Oct"]
    assert shown["ignored_sheets"] == ["Notes"]
    [row] = shown["payments"]
    assert row["row"] == 4 and row["status"] == "problem"  # 5 Oct 2026 is after "today"


# --------------------------------------------------------------------------- dates in files


def test_real_date_cells_and_money_formats(api: TestClient) -> None:
    make_student(api, name="Kabir Mehta")
    data = xlsx(
        (
            "Payments",
            [
                PAYMENT_HEAD,
                ["Kabir Mehta", "₹1,500", dt.datetime(2026, 5, 2), dt.date(2026, 5, 1), "cash"],
                ["Kabir Mehta", "1500/-", "2/5/2026", "May-26", "UPI"],
                ["Kabir Mehta", 1500.5, "2 May 2026", "may 2026", "UPI"],
            ],
        )
    )
    rows = by_row(preview(api, data)["payments"])
    assert [(r["amount_paise"], r["paid_on"], r["for_month"]) for r in rows.values()] == [
        (150000, "2026-05-02", "2026-05"),
        (150000, "2026-05-02", "2026-05"),
        (150050, "2026-05-02", "2026-05"),
    ]
    # Like row 2 (the 2 May is day first), but UPI, not cash: possibly the same payment.
    assert rows[3]["status"] == "possible_duplicate"
    assert rows[4]["status"] == "ready"


def test_the_template_reads_back(api: TestClient) -> None:
    for kind in ("students", "payments"):
        response = api.get("/api/import/template.xlsx", params={"kind": kind})
        assert response.status_code == 200
        assert f"scrappy-records-{kind}-template.xlsx" in response.headers["content-disposition"]
        book = load_workbook(io.BytesIO(response.content))
        assert book.sheetnames[1] == "How to fill this in"
        # Blank, so nothing to add, but its headings are recognized.
        shown = preview(api, response.content)
        assert shown["students"] == [] and shown["payments"] == []
        assert shown["sheets"] == [kind.capitalize()]


def test_pre_import_is_a_backup_reason() -> None:
    from app import backup

    assert "pre-import" in backup.REASONS
