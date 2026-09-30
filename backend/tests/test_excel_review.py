"""Uploads: the cases a careful review found (stale sheet sizes, hidden sheets, merged cells,
look-alike names, payments that only share a phone, possible duplicates), and speed."""

from __future__ import annotations

import io
import re
import time
import zipfile
from typing import Any

from fastapi.testclient import TestClient
from helpers import make_student, pay
from openpyxl import Workbook
from test_excel_import import by_row, commit, preview, preview_error, xlsx

STUDENT_HEAD = ["Name", "Phone", "Monthly fee", "Joined"]
PAYMENT_HEAD = ["Student", "Phone", "Amount", "Paid on", "For month", "Method", "Note"]


def _rewrite(data: bytes, member: str, change: Any) -> bytes:
    """The same .xlsx with one of its files changed."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == member:
                content = change(content)
            dst.writestr(item, content)
    return out.getvalue()


def test_a_stale_sheet_size_never_drops_rows(api: TestClient) -> None:
    rows = [STUDENT_HEAD, *[[f"Student {n}", None, 1000, "Jan 2026", "extra"] for n in range(8)]]
    data = xlsx(("Students", rows))
    # Say the sheet is only A1:C4, as some programs leave it after rows were added.
    stale = _rewrite(
        data,
        "xl/worksheets/sheet1.xml",
        lambda xml: re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="A1:C4"', xml),
    )
    assert b'ref="A1:C4"' in zipfile.ZipFile(io.BytesIO(stale)).read("xl/worksheets/sheet1.xml")
    shown = preview(api, stale)
    assert len(shown["students"]) == 8
    assert {s["joined_month"] for s in shown["students"]} == {"2026-01"}  # column D was read


def test_hidden_sheets_are_skipped_and_named(api: TestClient) -> None:
    book = Workbook()
    book.active.title = "Students"
    book.active.append(STUDENT_HEAD)
    book.active.append(["Kabir Mehta", None, 1500, None])
    secret = book.create_sheet("Old payments")
    secret.append(PAYMENT_HEAD)
    secret.append(["Kabir Mehta", None, 1500, "2026-05-02", "May 2026", "UPI", None])
    secret.sheet_state = "hidden"
    out = io.BytesIO()
    book.save(out)
    shown = preview(api, out.getvalue())
    assert shown["hidden_sheets"] == ["Old payments"]
    assert shown["payments"] == [] and len(shown["students"]) == 1


def test_merged_cells_count_on_every_row(api: TestClient) -> None:
    make_student(api, name="Kabir Mehta")
    book = Workbook()
    ws = book.active
    ws.title = "Payments"
    ws.append(PAYMENT_HEAD)
    ws.append(["Kabir Mehta", None, 1500, "2026-03-02", "Mar 2026", "UPI", None])
    ws.append([None, None, 1500, "2026-04-02", "Apr 2026", "UPI", None])
    ws.append([None, None, 1500, "2026-05-02", "May 2026", "UPI", None])
    ws.merge_cells("A2:A4")  # the name once, for three payments
    out = io.BytesIO()
    book.save(out)
    rows = by_row(preview(api, out.getvalue())["payments"])
    assert [rows[r]["student_text"] for r in (2, 3, 4)] == ["Kabir Mehta"] * 3
    assert {rows[r]["status"] for r in (2, 3, 4)} == {"ready"}


def test_a_payment_register_is_read_as_payments(api: TestClient) -> None:
    make_student(api, name="Kabir Mehta")
    data = xlsx(
        ("Sheet1", [["Name", "Date", "Fees", "Mode"], ["Kabir Mehta", "2/5/2026", 1500, "Cash"]])
    )
    shown = preview(api, data)
    assert shown["students"] == []
    [row] = shown["payments"]
    assert (row["status"], row["amount_paise"], row["method"]) == ("ready", 150000, "cash")
    # A students list with a fee is still a students list.
    listing = xlsx(("Sheet1", [["Name", "Fee", "Joined"], ["Diya Nair", 1200, "Jan 2026"]]))
    assert preview(api, listing)["students"][0]["status"] == "new"


def test_the_same_phone_under_another_name_needs_a_choice(api: TestClient) -> None:
    kabir = make_student(api, name="Kabir Mehta", phone="90000 00001")
    other = make_student(api, name="Diya Nair")
    data = xlsx(
        (
            "Payments",
            [
                PAYMENT_HEAD,
                ["Sunil Mehta", "90000 00001", 1500, "2026-05-02", "May 2026", "UPI", None],
                ["90000 00001", None, 1500, "2026-04-02", "Apr 2026", "UPI", None],
            ],
        )
    )
    rows = by_row(preview(api, data)["payments"])
    assert rows[2]["status"] == "needs_student"
    assert rows[2]["candidate_ids"][0] == kabir["id"]
    assert rows[2]["reason"].startswith("Same phone as Kabir Mehta, but the name is “Sunil Mehta”")
    assert other["id"] not in rows[2]["candidate_ids"]
    # Only a phone number, and it's his: nothing disagrees, so it's his.
    assert (rows[3]["status"], rows[3]["student_id"]) == ("ready", kabir["id"])


def test_near_miss_and_hyphenated_names(api: TestClient) -> None:
    make_student(api, name="Ananya Rao")
    make_student(api, name="Mohammed Khan")
    make_student(api, name="Mary-Jane Dsouza")
    make_student(api, name="Ria Das")
    data = xlsx(
        (
            "Students",
            [
                STUDENT_HEAD,
                ["Ananyaa Rao", None, 1500, None],
                ["Mohammad Khan", None, 1500, None],
                ["Mary Jane Dsouza", None, 1500, None],
                ["Riya Das", None, 1500, None],  # 7 letters, one apart: close
                ["Kiara Sethi", None, 1500, None],
            ],
        )
    )
    rows = by_row(preview(api, data)["students"])
    assert rows[2]["status"] == "similar"
    assert rows[2]["reason"] == "Name is very close to Ananya Rao"
    assert rows[3]["status"] == "similar"
    assert rows[4]["status"] == "exists"  # a hyphen separates words
    assert rows[5]["status"] == "similar"
    assert rows[6]["status"] == "new"


def test_brothers_and_sisters_sharing_a_phone_are_added(api: TestClient) -> None:
    data = xlsx(
        (
            "Students",
            [
                STUDENT_HEAD,
                ["Tara Iyer", "98765 43210", 1200, None],
                ["Meera Iyer", "98765 43210", 1200, None],
            ],
        )
    )
    shown = preview(api, data)
    rows = by_row(shown["students"])
    assert rows[3]["status"] == "similar" and rows[3]["add_by_default"] is True
    assert rows[3]["reason"] == "Same phone as Tara Iyer (row 2), perhaps a brother or sister"
    assert commit(api, shown)["students_added"] == 2


def test_possible_duplicates_and_add_anyway(api: TestClient) -> None:
    kabir = make_student(api, name="Kabir Mehta")
    pay(api, kabir["id"], "2026-05", paid_on="2026-05-02", method="cash", note="May")
    data = xlsx(
        (
            "Payments",
            [
                PAYMENT_HEAD,
                ["Kabir Mehta", None, 1500, "2026-05-02", "May 2026", "Cash", "May"],  # exact
                ["Kabir Mehta", None, 1200, "2026-05-03", "May 2026", "UPI", None],
                ["Kabir Mehta", None, 1200, "2026-05-03", "May 2026", "Cash", None],  # like 3
                ["Kabir Mehta", None, 1200, "2026-05-03", "May 2026", "UPI", None],  # same as 3
            ],
        )
    )
    shown = preview(api, data)
    rows = by_row(shown["payments"])
    assert [rows[r]["status"] for r in (2, 3, 4, 5)] == [
        "duplicate",
        "ready",
        "possible_duplicate",
        "duplicate",
    ]
    assert rows[2]["reason"] == "Already logged: ₹1,500 paid on 2 May 2026 for May 2026"
    assert rows[4]["reason"] == "Like row 3 in this file (the same amount, day and month)"
    assert rows[5]["reason"] == "Same as row 3 in this file"
    # Two instalments really were paid on the same day: Add anyway.
    result = commit(api, shown, payments={5: {"choice": "add"}})
    assert result["payments_added"] == 2
    assert len(api.get("/api/payments", params={"student_id": kabir["id"]}).json()) == 3


def test_a_possible_duplicate_of_a_payment_already_here(api: TestClient) -> None:
    kabir = make_student(api, name="Kabir Mehta")
    pay(api, kabir["id"], "2026-05", paid_on="2026-05-02", method="cash")
    data = xlsx(
        (
            "Payments",
            [PAYMENT_HEAD, ["Kabir Mehta", None, 1500, "2026-05-02", "May 2026", "UPI", None]],
        )
    )
    [row] = preview(api, data)["payments"]
    assert (row["status"], row["reason"]) == (
        "possible_duplicate",
        "Possibly already logged: ₹1,500 paid on 2 May 2026 for May 2026, by Cash",
    )


def test_plain_messages_for_long_names_and_two_amounts(api: TestClient) -> None:
    data = xlsx(
        (
            "Students",
            [
                STUDENT_HEAD,
                ["x" * 30000, None, 1500, None],
                ["Kabir Mehta", None, "₹500 700", None],
            ],
        )
    )
    rows = by_row(preview(api, data)["students"])
    assert rows[2]["reason"] == "Name is too long (max 200 characters)"
    assert len(rows[2]["name"]) == 200  # not echoed in full
    assert rows[3]["reason"].startswith("Monthly fee “₹500 700” looks like more than one amount")


def test_headings_not_found_says_where_it_looked(api: TestClient) -> None:
    rows = [[None]] * 10 + [["Name", "Monthly fee"], ["Kabir Mehta", 1500]]
    message = preview_error(api, xlsx(("Sheet1", rows)))
    assert "no headings in the first 10 rows" in message


def test_restoring_thousands_of_records_is_quick(api: TestClient) -> None:
    """3,000 students and 6,000 payments, previewed and added in a few seconds."""
    book = Workbook()
    ws = book.active
    ws.title = "Students"
    ws.append(["Name", "Phone", "Monthly fee", "Joined", "Student ID"])
    for n in range(3000):
        ws.append([f"Student {n:04d} Rao", f"9{n:09d}", 1000, "Jan 2026", n + 1])
    payments = book.create_sheet("Payments")
    payments.append(["Student", "Amount", "Paid on", "For month", "Method", "Student ID"])
    for n in range(6000):
        payments.append([f"Student {n % 3000:04d} Rao", 1000, "2026-02-03", "Feb 2026", "UPI",
                         n % 3000 + 1])  # fmt: skip
    out = io.BytesIO()
    book.save(out)
    started = time.monotonic()
    shown = preview(api, out.getvalue())
    result = commit(api, shown)
    took = time.monotonic() - started
    assert (result["students_added"], result["payments_added"]) == (3000, 6000)
    assert took < 30, took  # about 3 s on a laptop; generous for slow CI machines

    # A plain list (no Student IDs) of 3,000 names, each a letter away from one already here:
    # every one is compared, through the index, and flagged.
    rows: list[list[Any]] = [["Name", "Monthly fee"]]
    rows += [[f"Student {n:04d} Rau", 1000] for n in range(3000)]
    started = time.monotonic()
    shown = preview(api, xlsx(("Students", rows)))
    took = time.monotonic() - started
    assert {s["status"] for s in shown["students"]} == {"similar"}
    assert took < 30, took
