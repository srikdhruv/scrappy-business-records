"""Excel downloads: the Students and Payments lists as shown, everything in one workbook, and
blank templates for uploading.

Every sheet has bold, frozen headings and sensible column widths. Dates are real Excel dates
("5 Oct 2026"), months are real dates shown as "Oct 2026" (so they sort, and read back exactly
whatever the spreadsheet app does to them), and money is in rupees with a ₹ format. The
headings are ones the upload understands (`app/services/spreadsheet.py`), so any download can
be uploaded again: a Download everything file restores every record into an empty app.
"""

from __future__ import annotations

import datetime as dt
import io
import uuid
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models import PaymentMethod, Student, UnassignedPayment
from app.months import format_month, parse_month
from app.schemas import (
    BalanceStatus,
    ExportTemplateKind,
    PaymentRead,
    PaymentSort,
    SortOrder,
    StudentListFilter,
    StudentRead,
)
from app.services import payments as payment_service
from app.services import students as student_service
from app.services.spreadsheet import TEMPLATE_HELP_SHEET
from app.services.text import student_matches

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

ID_HEADING = "Student ID (for restoring)"

_HEAD_FONT = Font(bold=True)
_HEAD_FILL = PatternFill("solid", fgColor="FBEFD5")  # the app's cream-marigold
_ID_FONT = Font(color="808080")

_STATUS_WORDS = {
    BalanceStatus.owes: "Owes",
    BalanceStatus.credit: "Credit",
    BalanceStatus.up_to_date: "Up to date",
}
_METHOD_WORDS = {PaymentMethod.upi: "UPI", PaymentMethod.cash: "Cash", PaymentMethod.other: "Other"}


@dataclass(frozen=True)
class _Column:
    heading: str
    width: float
    kind: str = "text"  # text, phone, money, date, month, id


def _money(paise: int | None) -> tuple[Any, str]:
    if paise is None:
        return None, "General"
    if paise % 100 == 0:
        return paise // 100, '"₹"#,##0'
    return paise / 100, '"₹"#,##0.00'


def _put(ws: Worksheet, row: int, index: int, column: _Column, value: Any) -> None:
    fmt = None
    if value is None:
        pass
    elif column.kind == "money":
        value, fmt = _money(value)
    elif column.kind == "date":
        fmt = "d mmm yyyy"
    elif column.kind == "month":
        value = parse_month(value) if isinstance(value, str) else value
        fmt = "mmm yyyy"
    elif column.kind in ("text", "phone") and isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = ws.cell(row=row, column=index, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # "=..." stays text: never a formula
        if column.kind == "phone":
            fmt = "@"
    if column.kind == "id":
        cell.font = _ID_FONT
    if fmt:
        cell.number_format = fmt


def _add_sheet(
    book: Workbook, title: str, columns: Sequence[_Column], rows: Iterable[Sequence[Any]]
) -> Worksheet:
    ws = book.create_sheet(title)
    for i, column in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=column.heading)
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = column.width
    for r, values in enumerate(rows, start=2):
        for i, (column, value) in enumerate(zip(columns, values, strict=True), start=1):
            _put(ws, r, i, column, value)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 20
    return ws


def _save(book: Workbook) -> bytes:
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _new_book() -> Workbook:
    book = Workbook()
    book.remove(book.active)  # type: ignore[arg-type]
    return book


# --------------------------------------------------------------------------- students

_STUDENT_COLUMNS = (
    _Column("Name", 28),
    _Column("Phone", 16, "phone"),
    _Column("Parent/guardian", 24),
    _Column("Class/batch", 24),
    _Column("Monthly fee ₹ (current)", 14, "money"),
    _Column("Joined (month)", 12, "month"),
    _Column("Left (month)", 12, "month"),
    _Column("Status", 12),
    _Column("Owes ₹", 12, "money"),
    _Column("Notes", 40),
)


def _student_values(s: StudentRead) -> list[Any]:
    return [
        s.name,
        s.phone,
        s.guardian_name,
        s.batch_label,
        s.monthly_fee_paise,
        s.joined_month,
        s.left_month,
        _STATUS_WORDS.get(s.status, str(s.status).replace("_", " ").capitalize()),
        s.owed_paise,
        s.notes,
    ]


def students_shown(
    session: Session, status: StudentListFilter, q: str | None, current_month: dt.date
) -> list[StudentRead]:
    """The students the Students page shows for this tab and search (its `studentMatches`)."""
    rows = student_service.list_students(session, status, None, current_month)
    if q and q.strip():
        rows = [
            s
            for s in rows
            if student_matches(
                q,
                name=s.name,
                phone=s.phone,
                guardian_name=s.guardian_name,
                batch_label=s.batch_label,
            )
        ]
    return rows


def students_workbook(students: Sequence[StudentRead]) -> bytes:
    book = _new_book()
    _add_sheet(book, "Students", _STUDENT_COLUMNS, (_student_values(s) for s in students))
    return _save(book)


# --------------------------------------------------------------------------- payments

_PAYMENT_COLUMNS = (
    _Column("Student", 28),
    _Column("Phone", 16, "phone"),
    _Column("Amount ₹", 12, "money"),
    _Column("Paid on", 13, "date"),
    _Column("For month", 12, "month"),
    _Column("Method", 10),
    _Column("Note", 40),
)


def payments_shown(
    session: Session,
    *,
    student_id: int | None,
    month: str | None,
    q: str | None,
    method: PaymentMethod | None,
    sort: PaymentSort,
    order: SortOrder,
) -> list[PaymentRead]:
    """The payments the Payments page shows for these filters, in its order."""
    rows = payment_service.list_payments(
        session, student_id=student_id, month=month, q=q, sort=sort, order=order
    )
    return [p for p in rows if method is None or p.method is method]


def _phones(session: Session) -> dict[int, str | None]:
    return {sid: phone for sid, phone in session.execute(select(Student.id, Student.phone))}


def _payment_values(p: PaymentRead, phones: dict[int, str | None]) -> list[Any]:
    return [
        p.student_name,
        phones.get(p.student_id),
        p.amount_paise,
        p.paid_on,
        p.for_month,
        _METHOD_WORDS[p.method],
        p.note,
    ]


def payments_workbook(session: Session, payments: Sequence[PaymentRead]) -> bytes:
    phones = _phones(session)
    book = _new_book()
    _add_sheet(book, "Payments", _PAYMENT_COLUMNS, (_payment_values(p, phones) for p in payments))
    return _save(book)


# --------------------------------------------------------------------------- everything

_ID = _Column(ID_HEADING, 14, "id")
_FEE_COLUMNS = (
    _Column("Student", 28),
    _Column("Fee from (month)", 14, "month"),
    _Column("Monthly fee ₹", 14, "money"),
    _Column("Kind", 16),
    _ID,
)
_UNASSIGNED_COLUMNS = (
    _Column("Name as written", 28),
    _Column("Phone", 16, "phone"),
    _Column("Amount ₹", 12, "money"),
    _Column("Paid on", 13, "date"),
    _Column("For month", 12, "month"),
    _Column("Method", 10),
    _Column("Note", 40),
    _Column("Came from", 28),
)
_KIND_WORDS = {"fee": "Fee", "away": "Away (no fee)"}


def everything_workbook(session: Session, current_month: dt.date) -> bytes:
    """Students, Fee history, Payments and Unassigned payments, with each student's uid (the
    Student ID column), which links the sheets, and finds the same students again when the
    file is uploaded into this app or any other."""
    uids = ensure_uids(session)
    students = student_service.list_students(session, StudentListFilter.all, None, current_month)
    order = {s.id: i for i, s in enumerate(students)}
    phones = {s.id: s.phone for s in students}
    book = _new_book()
    _add_sheet(
        book,
        "Students",
        (*_STUDENT_COLUMNS, _ID),
        ([*_student_values(s), uids[s.id]] for s in students),
    )

    rows = session.scalars(select(Student).options(selectinload(Student.fee_changes))).all()
    fee_rows = sorted(
        (
            (
                order.get(s.id, 0),
                f.effective_month,
                s.name,
                f.amount_paise,
                f.kind.value,
                uids[s.id],
            )
            for s in rows
            for f in s.fee_changes
        ),
    )
    _add_sheet(
        book,
        "Fee history",
        _FEE_COLUMNS,
        (
            [name, format_month(month), amount, _KIND_WORDS[kind], sid]
            for _, month, name, amount, kind, sid in fee_rows
        ),
    )

    payments = payment_service.list_payments(session)
    _add_sheet(
        book,
        "Payments",
        (*_PAYMENT_COLUMNS, _ID),
        ([*_payment_values(p, phones), uids[p.student_id]] for p in payments),
    )

    unassigned = session.scalars(
        select(UnassignedPayment).order_by(UnassignedPayment.paid_on, UnassignedPayment.id)
    ).all()
    _add_sheet(
        book,
        "Unassigned payments",
        _UNASSIGNED_COLUMNS,
        (
            [
                u.student_text,
                u.phone,
                u.amount_paise,
                u.paid_on,
                format_month(u.for_month),
                _METHOD_WORDS[u.method],
                u.note,
                u.source,
            ]
            for u in unassigned
        ),
    )
    return _save(book)


def ensure_uids(session: Session) -> dict[int, str]:
    """Every student's uid (see `Student.uid`), giving one to anyone who hasn't got one yet."""
    rows = list(session.scalars(select(Student)))
    missing = [s for s in rows if not s.uid]
    for student in missing:
        student.uid = uuid.uuid4().hex
    if missing:
        session.commit()
    return {s.id: s.uid for s in rows if s.uid}


# --------------------------------------------------------------------------- templates

_TEMPLATE_HELP: dict[ExportTemplateKind, list[str]] = {
    ExportTemplateKind.students: [
        "One student per row, under the headings on the Students sheet. Keep the headings.",
        "Name and Monthly fee are needed. The rest can be left empty.",
        "Monthly fee: in rupees, like 1500 or ₹1,500.",
        "Joined (month): the first month they pay for, like Oct 2026. Empty means this month.",
        "Left (month): only for someone who has stopped coming: the last month they pay for.",
        "Phone: helps tell apart two students with the same name.",
        "Then, in Scrappy Records: Students → Upload Excel. You'll see what will be added "
        "before anything is saved.",
    ],
    ExportTemplateKind.payments: [
        "One payment per row, under the headings on the Payments sheet. Keep the headings.",
        "Student: their name as it is in Scrappy Records (or their phone number).",
        "Amount: in rupees, like 1500 or ₹1,500.",
        "Paid on: the day they paid, like 5 Oct 2026 or 05/10/2026 (day first).",
        "For month: the month it's for, like Oct 2026. Empty means the month it was paid in.",
        "Method: UPI, Cash or Other. Empty means Other.",
        "Then, in Scrappy Records: Payments → Upload Excel. You'll see what will be added "
        "before anything is saved. Names that don't match a student can be kept as "
        "unassigned payments and given to a student later.",
    ],
}

_TEMPLATE_COLUMNS: dict[ExportTemplateKind, tuple[str, Sequence[_Column]]] = {
    ExportTemplateKind.students: (
        "Students",
        (
            _Column("Name", 28),
            _Column("Phone", 16, "phone"),
            _Column("Parent/guardian", 24),
            _Column("Class/batch", 24),
            _Column("Monthly fee ₹", 14, "money"),
            _Column("Joined (month)", 14, "month"),
            _Column("Left (month)", 14, "month"),
            _Column("Notes", 40),
        ),
    ),
    ExportTemplateKind.payments: ("Payments", _PAYMENT_COLUMNS),
}


def template_workbook(kind: ExportTemplateKind) -> bytes:
    """A blank sheet with the headings, and a sheet saying how to fill it in (which the upload
    skips)."""
    title, columns = _TEMPLATE_COLUMNS[kind]
    book = _new_book()
    ws = _add_sheet(book, title, columns, [])
    formats: dict[str, Callable[[], str]] = {
        "money": lambda: '"₹"#,##0',
        "date": lambda: "d mmm yyyy",
        "month": lambda: "mmm yyyy",
        "phone": lambda: "@",
    }
    for i, column in enumerate(columns, start=1):
        if column.kind in formats:  # typed-in values get the right look too
            for r in range(2, 202):
                ws.cell(row=r, column=i).number_format = formats[column.kind]()
    help_sheet = book.create_sheet(TEMPLATE_HELP_SHEET)
    help_sheet.column_dimensions["A"].width = 110
    help_sheet.cell(row=1, column=1, value="How to fill this in").font = Font(bold=True, size=13)
    for r, line in enumerate(_TEMPLATE_HELP[kind], start=3):
        help_sheet.cell(row=r, column=1, value=f"• {line}")
    return _save(book)


def filename(what: str, today: dt.date) -> str:
    return f"scrappy-records-{what}-{today.isoformat()}.xlsx"
