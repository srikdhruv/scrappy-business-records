"""The monthly report as an Excel file (`GET /api/report.xlsx`): the same rows and totals as
the Report screen, from `app.services.report.get_report`.

A title row ("Scrappy Records — Fees report, October 2026"), the date, then the headings
(bold, frozen with the Student column) and one row per student, then a bold totals row. Money is
in rupees with a ₹ number format (a nil amount shows as "—" but is still the number 0, so sums
work). The sheet prints on A4 landscape, one page wide.

Kept small and on its own on purpose. It follows the same approach as the Excel downloads
(`openpyxl`, the same ₹ formats, heading colour and text safety), so the two can share one
helper module once both are in.
"""

from __future__ import annotations

import datetime as dt
import io
from collections.abc import Sequence
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.months import parse_month
from app.schemas import CreditSource, ExtraSent, ReportResponse, ReportRow, ReportStatus

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

STATUS_WORDS = {
    ReportStatus.paid: "Paid",
    ReportStatus.paid_with_credit: "Paid with credit",
    ReportStatus.partial: "Partial",
    ReportStatus.unpaid: "Unpaid",
    ReportStatus.no_fee: "No fee",
    ReportStatus.not_due_yet: "Not due yet",
    ReportStatus.left: "Left",
}
"""The words on the Report screen (`frontend/src/lib/report.ts`)."""

_HEAD_FONT = Font(bold=True)
_HEAD_FILL = PatternFill("solid", fgColor="FBEFD5")  # the app's cream-marigold
_TITLE_FONT = Font(bold=True, size=14)
_TOTAL_BORDER = Border(top=Side(style="thin"))
# Whole rupees, or with paise when there are any; nothing shows as "—" (still the number 0).
_RUPEES = '"₹"#,##0;-"₹"#,##0;"—"'
_RUPEES_PAISE = '"₹"#,##0.00;-"₹"#,##0.00;"—"'

HEADER_ROW = 4
"""Row 1 is the title, row 2 the date, row 3 is empty."""


@dataclass(frozen=True)
class _Column:
    heading: str
    width: float
    money: bool = False


COLUMNS: tuple[_Column, ...] = (
    _Column("Student", 26),
    _Column("Class/batch", 20),
    _Column("Phone", 15),
    _Column("Fee ₹", 11, money=True),
    _Column("Paid for this month ₹ (logged)", 14, money=True),
    _Column("Covered by credit ₹", 13, money=True),
    _Column("Credit came from", 34),
    _Column("Extra sent elsewhere ₹", 13, money=True),
    _Column("Extra went to", 28),
    _Column("Short ₹", 11, money=True),
    _Column("Status", 16),
    _Column("Owed from earlier months ₹", 14, money=True),
    _Column("Earlier months owed", 22),
    _Column("Total owed now ₹", 13, money=True),
    _Column("Credit ₹", 11, money=True),
    _Column("Paid ahead ₹", 12, money=True),
)


def rupees(paise: int) -> str:
    """₹ in Indian grouping, as on screen: 15000000 -> "₹1,50,000"; 150050 -> "₹1,500.50"."""
    whole, part = divmod(abs(paise), 100)
    digits = str(whole)
    if len(digits) > 3:
        head, tail = digits[:-3], digits[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        digits = ",".join([head, *groups, tail])
    sign = "-" if paise < 0 else ""
    return f"{sign}₹{digits}" + (f".{part:02d}" if part else "")


def _date(value: dt.date) -> str:
    return f"{value.day} {value:%b %Y}"


def _month_short(month: str) -> str:
    return f"{parse_month(month):%b %Y}"


def month_long(month: str) -> str:
    return f"{parse_month(month):%B %Y}"


def credit_from_text(sources: Sequence[CreditSource]) -> str:
    """ "₹1,500 from the 5 Sep 2026 payment (for Sep 2026)", one per source."""
    return "; ".join(
        f"{rupees(s.amount_paise)} from the {_date(s.paid_on)} payment "
        f"(for {_month_short(s.for_month)})"
        for s in sources
    )


def extra_went_text(sent: Sequence[ExtraSent], unused_paise: int) -> str:
    """ "₹1,500 → Aug 2026", one per month, and "₹500 kept as credit"."""
    parts = [f"{rupees(e.amount_paise)} → {_month_short(e.to_month)}" for e in sent]
    if unused_paise:
        parts.append(f"{rupees(unused_paise)} kept as credit")
    return "; ".join(parts)


def _values(r: ReportRow) -> list[Any]:
    return [
        r.student_name,
        r.batch_label,
        r.phone,
        r.fee_paise,
        r.paid_paise,
        r.covered_by_credit_paise,
        credit_from_text(r.credit_sources),
        r.extra_sent_paise,
        extra_went_text(r.extra_sent, r.extra_unused_paise),
        r.short_paise,
        STATUS_WORDS[r.status],
        r.owed_before_paise,
        ", ".join(_month_short(m) for m in r.owed_before_months),
        r.owed_now_paise,
        r.credit_paise,
        r.paid_ahead_paise,
    ]


def _put(ws: Worksheet, row: int, col: int, column: _Column, value: Any) -> None:
    fmt = None
    if column.money and isinstance(value, int):
        fmt = _RUPEES if value % 100 == 0 else _RUPEES_PAISE
        value = value // 100 if value % 100 == 0 else value / 100
    elif isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value) or None
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # a name starting with "=" stays text: never a formula
    if fmt:
        cell.number_format = fmt


def filename(month: str) -> str:
    return f"scrappy-records-report-{month}.xlsx"


def title(month: str) -> str:
    return f"Scrappy Records — Fees report, {month_long(month)}"


def workbook(report: ReportResponse) -> bytes:
    book = Workbook()
    ws = book.active
    assert ws is not None
    ws.title = f"Report {_month_short(report.month)}"
    last_col = get_column_letter(len(COLUMNS))

    ws["A1"] = title(report.month)
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"As of {_date(report.today)}"
    if report.month > report.current_month:
        ws["A2"] = (
            f"As of {_date(report.today)}. {month_long(report.month)} isn't due yet: what's "
            "paid for it is paid ahead."
        )

    for i, column in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=i, value=column.heading)
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = column.width
    ws.row_dimensions[HEADER_ROW].height = 32

    row = HEADER_ROW
    for r in report.rows:
        row += 1
        for i, (column, value) in enumerate(zip(COLUMNS, _values(r), strict=True), start=1):
            _put(ws, row, i, column, value)

    t = report.totals
    row += 1
    count = f"{t.student_count} student" + ("" if t.student_count == 1 else "s")
    totals: dict[int, Any] = {
        1: f"Total ({count})",
        4: t.fee_paise,
        5: t.paid_paise,
        6: t.covered_by_credit_paise,
        8: t.extra_sent_paise,
        10: t.short_paise,
        11: f"{t.not_fully_paid_count} of {t.active_student_count} not fully paid",
        12: t.owed_before_paise,
        14: t.owed_now_paise,
        15: t.credit_paise,
        16: t.paid_ahead_paise,
    }
    for i, column in enumerate(COLUMNS, start=1):
        _put(ws, row, i, column, totals.get(i))
        cell = ws.cell(row=row, column=i)
        cell.font = _HEAD_FONT
        cell.border = _TOTAL_BORDER

    ws.freeze_panes = f"B{HEADER_ROW + 1}"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col}{max(row - 1, HEADER_ROW)}"
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
