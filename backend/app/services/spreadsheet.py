"""Reading an uploaded Excel file forgivingly: finding the sheets and columns, and turning cells
into the app's values (paise, dates, "YYYY-MM" months).

Only the shape of the file is decided here. Whether a row is new, already here or a problem is
`app/services/imports.py`. Every failure is a `BadFile` (the whole file) or a `CellError` (one
cell), each with a plain-words message the owner can act on.

What is understood (case, spaces, punctuation and anything in brackets in a heading don't
matter, so the app's own downloads read back as they are):

- Dates: real date cells, "5 Oct 2026", "05/10/2026" (day first, as in India), "5-10-26",
  "2026-10-05".
- Months: real date cells (their month), "Oct 2026", "October 2026", "Oct-26", "2026-10",
  "10/2026".
- Money: numbers, "₹1,500", "1500/-", "Rs. 1,500.00", "1,50,000".
"""

from __future__ import annotations

import datetime as dt
import io
import math
import re
import warnings
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from enum import StrEnum
from typing import Any

from app.models import FeeKind, PaymentMethod
from app.months import format_month

MAX_FILE_BYTES = 5 * 1024 * 1024
"""The largest file accepted (5 MB): years of records are far smaller."""
MAX_ROWS = 5000
"""The most rows read from one sheet of a list someone made."""
MAX_ROWS_WITH_IDS = 300_000
"""The most rows in a sheet with the app's Student ID column (a Download everything file, which
is as big as the records are), and in the whole file. A 5 MB file can't really hold more: the
file size is the real limit."""
_MAX_UNZIPPED_BYTES = 80 * 1024 * 1024  # an .xlsx is a zip; refuse one that balloons
_MAX_ZIP_ENTRIES = 2000
_HEADER_SEARCH_ROWS = 10  # the headings may be below a title or a blank row or two

TEMPLATE_HELP_SHEET = "How to fill this in"


class BadFile(Exception):
    """The file as a whole can't be used. `message` is shown to the owner as it is."""

    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


class CellError(ValueError):
    """One cell couldn't be read. The message is plain words, without the column name."""


# --------------------------------------------------------------------------- columns


class Col(StrEnum):
    name = "name"
    phone = "phone"
    guardian = "guardian"
    batch = "batch"
    fee = "fee"
    joined = "joined"
    left = "left"
    notes = "notes"
    ref = "ref"
    student = "student"
    amount = "amount"
    paid_on = "paid_on"
    for_month = "for_month"
    method = "method"
    note = "note"
    fee_from = "fee_from"
    kind = "kind"
    source = "source"
    ignored = "ignored"


def normalize_heading(value: object) -> str:
    """ "Monthly fee ₹ (current)" -> "monthly fee"; "Parent / Guardian" -> "parent/guardian"."""
    text = str(value or "").lower()
    text = re.sub(r"\([^)]*\)", " ", text)  # anything in brackets
    text = re.sub(r"₹|\brs\b\.?|\binr\b", " ", text)
    text = re.sub(r"\s*/\s*", "/", text)
    text = re.sub(r"[^a-z0-9/ ]+", " ", text)
    return " ".join(text.split())


_STUDENT_HEADINGS = {
    Col.name: ("name", "student", "student name", "full name", "name of student"),
    Col.phone: (
        "phone",
        "phone number",
        "phone no",
        "mobile",
        "mobile number",
        "mobile no",
        "contact",
        "contact number",
        "whatsapp",
    ),
    Col.guardian: (
        "parent/guardian",
        "parent",
        "guardian",
        "parent name",
        "guardian name",
        "parent or guardian",
        "parents",
    ),
    Col.batch: ("class/batch", "class", "batch", "class or batch", "batch/class", "group"),
    Col.fee: ("monthly fee", "fee", "fees", "monthly fees", "fee per month"),
    Col.joined: (
        "joined",
        "joined in",
        "joined month",
        "joining month",
        "start month",
        "started",
        "joined on",
        "date of joining",
    ),
    Col.left: ("left", "left in", "left month", "left on"),
    Col.notes: ("notes", "note", "remarks", "comments"),
    Col.ref: ("student id", "id"),
    Col.ignored: ("status", "owes", "owed"),
}

_STRICT_AMOUNT = ("amount", "amount paid", "paid", "payment", "amount received")

_PAYMENT_HEADINGS = {
    Col.student: ("student", "name", "student name", "name as written", "paid by"),
    Col.phone: _STUDENT_HEADINGS[Col.phone],
    Col.amount: (*_STRICT_AMOUNT, "fee", "fees", "fee paid", "fees paid", "fees received"),
    Col.paid_on: ("paid on", "date", "payment date", "paid date", "date paid", "received on"),
    Col.for_month: ("for month", "month", "fee month", "for the month"),
    Col.method: ("method", "mode", "payment method", "payment mode", "how they paid"),
    Col.note: ("note", "notes", "remarks", "comments"),
    Col.ref: ("student id",),
    Col.source: ("came from", "source"),
}

_FEE_HEADINGS = {
    Col.student: ("student", "name"),
    Col.fee_from: ("fee from", "from", "effective month", "from month", "starts"),
    Col.fee: ("monthly fee", "fee", "amount"),
    Col.kind: ("kind", "type"),
    Col.ref: ("student id",),
}


class SheetKind(StrEnum):
    students = "students"
    payments = "payments"
    fee_history = "fee_history"
    unassigned = "unassigned"


_HEADINGS = {
    SheetKind.students: _STUDENT_HEADINGS,
    SheetKind.payments: _PAYMENT_HEADINGS,
    SheetKind.unassigned: _PAYMENT_HEADINGS,
    SheetKind.fee_history: _FEE_HEADINGS,
}

# The app's own sheet names, which decide a sheet's kind before its headings do.
_TITLES = {
    "students": SheetKind.students,
    "fee history": SheetKind.fee_history,
    "payments": SheetKind.payments,
    "unassigned payments": SheetKind.unassigned,
}


def _columns(kind: SheetKind, headings: Iterable[object]) -> dict[Col, int]:
    """Column index for each column recognized, by its heading. The first match wins."""
    table = _HEADINGS[kind]
    found: dict[Col, int] = {}
    for index, raw in enumerate(headings):
        heading = normalize_heading(raw)
        if not heading:
            continue
        for col, names in table.items():
            if heading in names and col not in found:
                found[col] = index
                break
    found.pop(Col.ignored, None)
    return found


def _complete(kind: SheetKind, cols: dict[Col, int]) -> bool:
    """Enough columns to be a sheet of this kind."""
    if kind is SheetKind.students:
        return Col.name in cols and (Col.fee in cols or Col.joined in cols)
    if kind in (SheetKind.payments, SheetKind.unassigned):
        return (Col.student in cols or Col.phone in cols) and Col.amount in cols
    return Col.fee_from in cols and Col.fee in cols and Col.ref in cols


# --------------------------------------------------------------------------- reading


@dataclass
class Row:
    number: int  # the row number in the spreadsheet, as the owner sees it
    cells: dict[Col, Any]

    def get(self, col: Col) -> Any:
        return self.cells.get(col)


@dataclass
class Sheet:
    title: str
    kind: SheetKind
    columns: dict[Col, int]
    rows: list[Row] = field(default_factory=list)


@dataclass
class Workbook:
    sheets: list[Sheet]
    ignored: list[str]  # titles of sheets that weren't read
    hidden: list[str] = field(default_factory=list)  # hidden sheets, never read

    def of(self, kind: SheetKind) -> list[Sheet]:
        return [s for s in self.sheets if s.kind is kind]


def _check_container(data: bytes) -> None:
    if not data:
        raise BadFile("The file is empty.")
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise BadFile(
            "This is an older kind of Excel file (.xls). Open it in Excel, choose File → "
            "Save As → Excel Workbook (.xlsx), and upload that instead."
        )
    not_xlsx = BadFile(
        "This isn't an Excel file (.xlsx). Choose a file saved from Excel, Google Sheets or "
        "LibreOffice as an Excel workbook, or one downloaded from Scrappy Records."
    )
    if data[:2] != b"PK":
        raise not_xlsx
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = {e.filename for e in entries}
    except (zipfile.BadZipFile, ValueError, OSError, EOFError):
        raise not_xlsx from None
    if "[Content_Types].xml" not in names or not any(n.startswith("xl/") for n in names):
        raise not_xlsx
    if len(entries) > _MAX_ZIP_ENTRIES or sum(e.file_size for e in entries) > _MAX_UNZIPPED_BYTES:
        raise BadFile("This file is too big to read. Upload one with fewer rows.")


_MERGE_COLUMNS = 64  # no sheet the app reads is wider; a merge wider than this is cut
_MERGES = 20_000  # more merged ranges than this and the rest are left as they are


def _merged(ws: Any) -> list[tuple[int, int, int, int]]:
    """A sheet's merged ranges as (first row, last row, first column, last column), sorted by
    first row, read straight from the sheet's XML (read-only sheets don't keep them). Only
    the first `_MERGE_COLUMNS` columns count, so one huge merge in a tiny file costs nothing.
    Empty if unreadable."""
    from xml.etree.ElementTree import iterparse

    from openpyxl.utils.cell import range_boundaries

    ranges: list[tuple[int, int, int, int]] = []
    try:
        with ws._get_source() as source:  # the sheet's XML inside the .xlsx
            for _, element in iterparse(source):
                if element.tag.endswith("}mergeCell") and len(ranges) < _MERGES:
                    first_col, first_row, last_col, last_row = range_boundaries(element.get("ref"))
                    if first_col <= _MERGE_COLUMNS:
                        last_col = min(last_col, _MERGE_COLUMNS)
                        ranges.append((first_row, last_row, first_col, last_col))
                element.clear()
    except Exception:  # never let an odd merge stop the upload; the cells just stay empty
        return []
    return sorted(ranges)


def _fill_merged(
    rows: Iterable[tuple[Any, ...]], merged: list[tuple[int, int, int, int]]
) -> Iterator[tuple[Any, ...]]:
    """Rows with each merged range's value copied into all of its cells, as it looks in Excel
    (a name merged down three rows is the name on each of them). Only rows the sheet really
    has are read, so a range reaching row 1,000,000 costs no more than the rows present."""
    if not merged:
        yield from rows
        return
    upcoming = list(merged)
    upcoming.reverse()  # pop() gives the next range to start
    active: list[tuple[int, int, int, int, Any]] = []  # ranges covering the current row, with value
    for number, values in enumerate(rows, start=1):
        active = [r for r in active if r[1] >= number]
        starting: list[tuple[int, int, int, int]] = []
        while upcoming and upcoming[-1][0] <= number:
            first = upcoming.pop()
            if first[1] >= number:
                starting.append(first)
        if not active and not starting:
            yield values
            continue
        filled = list(values)
        for first_row, last_row, first_col, last_col in starting:
            if len(filled) < first_col:
                filled += [None] * (first_col - len(filled))
            # A range that started on a row that wasn't read (it was empty) has no value.
            value = filled[first_col - 1] if first_row == number else None
            active.append((first_row, last_row, first_col, last_col, value))
        for _, _, first_col, last_col, value in active:
            if len(filled) < last_col:
                filled += [None] * (last_col - len(filled))
            for c in range(first_col, last_col + 1):
                filled[c - 1] = value
        yield tuple(filled)


def read_workbook(data: bytes) -> Workbook:
    """Find the student, payment and fee-history sheets in an .xlsx file, and read their rows.

    Hidden sheets are left out (and listed). Merged cells count as their value in every cell.
    Raises `BadFile` for anything that isn't a readable .xlsx, too many rows, or a file with
    nothing the app recognizes.
    """
    if len(data) > MAX_FILE_BYTES:
        raise BadFile("This file is bigger than 5 MB. Upload a smaller one.")
    _check_container(data)
    try:
        import openpyxl  # imported here: only uploads and downloads need it

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # e.g. "Workbook contains no default style"
            book = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=True, keep_links=False
            )
    except Exception:  # openpyxl raises many kinds of errors for a damaged file
        raise BadFile(
            "Couldn't open this Excel file. It may be damaged, or protected with a password. "
            "Open it in Excel, save it again as an Excel workbook (.xlsx), and upload that."
        ) from None

    sheets: list[Sheet] = []
    ignored: list[str] = []
    hidden: list[str] = []
    total = 0
    try:
        for ws in book.worksheets:
            title = str(ws.title)
            if title == TEMPLATE_HELP_SHEET:
                continue
            if getattr(ws, "sheet_state", "visible") != "visible":
                hidden.append(title)
                continue
            # The size a file claims for a sheet can be out of date; never trust it, or rows
            # past it would be silently left out.
            ws.reset_dimensions()
            rows = _fill_merged(ws.iter_rows(values_only=True), _merged(ws))
            sheet = _read_sheet(title, rows)
            # One sheet of each kind: a second one is left out, and said so.
            if sheet is None or any(s.kind is sheet.kind for s in sheets):
                ignored.append(title)
                continue
            sheets.append(sheet)
            total += len(sheet.rows)
            if total > MAX_ROWS_WITH_IDS:
                raise BadFile(
                    f"This file has more than {MAX_ROWS_WITH_IDS:,} rows. Split it into smaller "
                    "files and upload them one at a time."
                )
    except BadFile:
        raise
    except Exception:
        raise BadFile(
            "Couldn't read this Excel file. Open it in Excel, save it again as an Excel "
            "workbook (.xlsx), and upload that."
        ) from None
    finally:
        book.close()

    readable = (SheetKind.students, SheetKind.payments, SheetKind.unassigned)
    if not any(s.kind in readable for s in sheets):
        raise BadFile(
            "Couldn't find students or payments in this file: no headings in the first 10 rows "
            "of any sheet. It needs headings such as Name and Monthly fee for students, or "
            "Student, Amount and Paid on for payments. Download a blank template to see how."
        )
    if not any(s.kind is SheetKind.students for s in sheets):
        # A fee history means nothing without its students.
        ignored += [s.title for s in sheets if s.kind is SheetKind.fee_history]
        sheets = [s for s in sheets if s.kind is not SheetKind.fee_history]
    return Workbook(sheets=sheets, ignored=ignored, hidden=hidden)


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _read_sheet(title: str, rows: Iterable[tuple[Any, ...]]) -> Sheet | None:
    """The sheet's kind, columns and non-empty rows, or None if it isn't one the app reads."""
    by_title = _TITLES.get(normalize_heading(title))
    sheet: Sheet | None = None
    limit = MAX_ROWS
    for number, values in enumerate(rows, start=1):
        if sheet is None:
            if number > _HEADER_SEARCH_ROWS:
                return None
            if all(_blank(v) for v in values):
                continue
            kinds = [by_title] if by_title else list(_guess_order(values))
            for kind in kinds:
                cols = _columns(kind, values)
                if _complete(kind, cols):
                    sheet = Sheet(title=title, kind=kind, columns=cols)
                    # The app's own Download everything sheets carry Student IDs: those can be
                    # as big as the records are.
                    limit = MAX_ROWS_WITH_IDS if Col.ref in cols else MAX_ROWS
                    break
            continue
        if all(_blank(v) for v in values):
            continue
        if len(sheet.rows) >= limit:
            raise BadFile(
                f"The sheet “{title}” has more than {limit:,} rows. Split it into smaller "
                "files and upload them one at a time."
            )
        cells = {col: values[i] if i < len(values) else None for col, i in sheet.columns.items()}
        sheet.rows.append(Row(number=number, cells=cells))
    return sheet


def _guess_order(values: tuple[Any, ...]) -> list[SheetKind]:
    """Which kinds to try for a sheet with a name of its own. Payments first if it has an
    amount, a date or a payment-method column (a payments list has names too, and "Name | Date
    | Fees | Mode" is a list of payments); else students first; fee history when it says so."""
    headings = {normalize_heading(v) for v in values}
    order = [SheetKind.students, SheetKind.payments, SheetKind.fee_history]
    # An amount or a payment method means payments. A date alone is weaker: a list of students
    # can have a date (of joining) too, so it counts only without a phone column, which a
    # list of payments rarely has ("Name | Mobile | Fee | Date" is students; "Name | Date |
    # Fees" is payments).
    paymentish = set(_STRICT_AMOUNT) | set(_PAYMENT_HEADINGS[Col.method])
    if not headings & set(_PAYMENT_HEADINGS[Col.phone]):
        paymentish |= set(_PAYMENT_HEADINGS[Col.paid_on])
    if headings & paymentish:
        order = [SheetKind.payments, SheetKind.students, SheetKind.fee_history]
    if headings & set(_FEE_HEADINGS[Col.fee_from]) and headings & set(_FEE_HEADINGS[Col.kind]):
        order = [SheetKind.fee_history, *[k for k in order if k is not SheetKind.fee_history]]
    return order


# --------------------------------------------------------------------------- cells


def text(value: object) -> str | None:
    """A cell as text, or None if blank. Whole numbers lose their ".0" (a phone typed as a
    number reads 9876543210, not 9876543210.0)."""
    if value is None or isinstance(value, bool):
        return None if value is None else str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    result = str(value).strip()
    return result or None


_MONEY_NOISE = re.compile(r"₹|inr|rs\.?|/-|/=|,", re.IGNORECASE)
_MONEY_RE = re.compile(r"^\d+(\.\d{1,2})?$")


def money(value: object) -> int | None:
    """Rupees in a cell -> paise, or None if blank. 1500, "₹1,500", "1500/-", "Rs. 1,50,000.00".
    Raises `CellError` for anything else (negative amounts, text, more than 2 decimals)."""
    if _blank(value):
        return None
    if isinstance(value, bool):
        raise CellError("isn't an amount")
    if isinstance(value, int | float):
        if not math.isfinite(value):
            raise CellError("isn't an amount")
        if value < 0:
            raise CellError("can't be less than ₹0")
        paise = round(value * 100)
        if abs(value * 100 - paise) > 1e-6:
            raise CellError(f"“{value}” has more than 2 decimal places")
        return int(paise)
    raw = str(value).strip()
    if re.search(r"\d[\s,]*\s[\s,]*\d", raw):
        # "₹500 700" is two amounts, or a typo: never ₹5,00,700.
        raise CellError(f"“{raw[:40]}” looks like more than one amount. Write one amount per row")
    cleaned = _MONEY_NOISE.sub("", re.sub(r"\s", "", raw))
    if cleaned.startswith("-"):
        raise CellError("can't be less than ₹0")
    if not _MONEY_RE.match(cleaned):
        raise CellError(f"“{raw[:40]}” isn't an amount. Write it like 1500 or ₹1,500")
    rupees, _, paise = cleaned.partition(".")
    return int(rupees) * 100 + int((paise + "00")[:2])


_MONTH_NAMES = {
    name: i
    for i, names in enumerate(
        (
            ("jan", "january"),
            ("feb", "february"),
            ("mar", "march"),
            ("apr", "april"),
            ("may",),
            ("jun", "june"),
            ("jul", "july"),
            ("aug", "august"),
            ("sep", "sept", "september"),
            ("oct", "october"),
            ("nov", "november"),
            ("dec", "december"),
        ),
        start=1,
    )
    for name in names
}
_EXCEL_EPOCH = dt.date(1899, 12, 30)
_SERIAL_RANGE = (36526, 73050)  # 1 Jan 2000 .. 31 Dec 2099 as Excel day numbers


def _year(raw: str) -> int:
    year = int(raw)
    return 2000 + year if len(raw) <= 2 else year


def _serial(value: float) -> dt.date | None:
    if _SERIAL_RANGE[0] <= value <= _SERIAL_RANGE[1]:
        return _EXCEL_EPOCH + dt.timedelta(days=int(value))
    return None


def _make_date(year: int, month: int, day: int) -> dt.date | None:
    try:
        return dt.date(year, month, day)
    except ValueError:
        return None


# A time after a date ("05/10/2026 10:30", "5 Oct 2026 4:15 pm") is ignored.
_TIME = r"(?:[\s,t]+\d{1,2}[:.]\d{2}(?:[:.]\d{2})?\s*(?:am|pm)?)?"
_ISO_DATE = re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:[ t].*)?$", re.IGNORECASE)
_DMY = re.compile(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2}|\d{4})" + _TIME + "$")
_D_MON_Y = re.compile(
    r"^(\d{1,2})(?:st|nd|rd|th)?[\s\-/.,]*([a-z]+)[\s\-/.,']*(\d{2}|\d{4})" + _TIME + "$"
)
_MON_D_Y = re.compile(
    r"^([a-z]+)[\s\-/.]*(\d{1,2})(?:st|nd|rd|th)?[\s,\-/.']+(\d{2}|\d{4})" + _TIME + "$"
)
_SERIAL_TEXT = re.compile(r"^\d{5}(\.\d+)?$")


def date(value: object) -> dt.date | None:
    """A date cell -> a date, or None if blank. Text is read day first, as in India:
    "05/10/2026" is 5 October 2026. Raises `CellError` if it can't be read."""
    if _blank(value):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, int | float) and not isinstance(value, bool):
        found = _serial(value) if math.isfinite(value) else None
        if found:
            return found
        raise CellError(f"“{text(value)}” isn't a date. Write it like 5 Oct 2026 or 05/10/2026")
    raw = str(value).strip()
    s = raw.lower()
    found: dt.date | None = None
    if _SERIAL_TEXT.match(s):  # an Excel day number that became text, e.g. "46300"
        found = _serial(float(s))
    elif m := _ISO_DATE.match(s):
        found = _make_date(int(m[1]), int(m[2]), int(m[3]))
    elif m := _DMY.match(s):
        found = _make_date(_year(m[3]), int(m[2]), int(m[1]))
    elif (m := _D_MON_Y.match(s)) and m[2] in _MONTH_NAMES:
        found = _make_date(_year(m[3]), _MONTH_NAMES[m[2]], int(m[1]))
    elif (m := _MON_D_Y.match(s)) and m[1] in _MONTH_NAMES:
        found = _make_date(_year(m[3]), _MONTH_NAMES[m[1]], int(m[2]))
    if found is None:
        raise CellError(f"“{raw[:40]}” isn't a date. Write it like 5 Oct 2026 or 05/10/2026")
    return found


_MON_Y = re.compile(r"^([a-z]+)[\s\-/.,']*(\d{2}|\d{4})$")
_Y_M = re.compile(r"^(\d{4})[-/.](\d{1,2})$")
_M_Y = re.compile(r"^(\d{1,2})[-/.](\d{4})$")
_M_YY = re.compile(r"^(\d{1,2})[-/.](\d{2})$")  # "10/26": October 2026


def month(value: object) -> str | None:
    """A month cell -> "YYYY-MM", or None if blank. A date means its month. Raises
    `CellError` if it can't be read, or is outside 2000-2099."""
    if _blank(value):
        return None
    problem = CellError(f"“{(text(value) or '')[:40]}” isn't a month. Write it like Oct 2026")
    found: tuple[int, int] | None = None
    if isinstance(value, dt.date):  # datetimes are dates too
        found = (value.year, value.month)
    elif isinstance(value, int | float) and not isinstance(value, bool):
        serial = _serial(value) if math.isfinite(value) else None
        if serial is None:
            raise problem
        found = (serial.year, serial.month)
    else:
        s = str(value).strip().lower()
        if (m := _MON_Y.match(s)) and m[1] in _MONTH_NAMES:
            found = (_year(m[2]), _MONTH_NAMES[m[1]])
        elif m := _Y_M.match(s):
            found = (int(m[1]), int(m[2]))
        elif m := _M_Y.match(s):
            found = (int(m[2]), int(m[1]))
        elif m := _M_YY.match(s):
            found = (_year(m[2]), int(m[1]))
        else:
            try:
                d = date(value)
            except CellError:
                raise problem from None
            found = (d.year, d.month) if d else None
    if found is None or not 1 <= found[1] <= 12:
        raise problem
    year, mon = found
    if not 2000 <= year <= 2099:
        raise CellError(f"{dt.date(2000, mon, 1):%B} {year} is outside the years 2000 to 2099")
    return format_month(dt.date(year, mon, 1))


_UPI = {"upi", "gpay", "google pay", "googlepay", "phonepe", "phone pe", "paytm", "bhim"}
_CASH = {"cash"}


def method(value: object) -> PaymentMethod:
    """UPI (and GPay, PhonePe, Paytm, BHIM), Cash, or Other for anything else, blank included."""
    s = " ".join((text(value) or "").lower().split())
    if s in _UPI:
        return PaymentMethod.upi
    if s in _CASH:
        return PaymentMethod.cash
    return PaymentMethod.other


def fee_kind(value: object) -> FeeKind:
    s = normalize_heading(value)
    if s in ("", "fee"):
        return FeeKind.fee
    if s.startswith("away"):
        return FeeKind.away
    raise CellError(f"“{(text(value) or '')[:40]}” should be Fee or Away")
